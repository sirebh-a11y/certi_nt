"""Local-only integration tests for a selected persistent DDT share."""
from contextlib import ExitStack, contextmanager
from concurrent.futures import ThreadPoolExecutor, TimeoutError
from datetime import datetime, timezone
from decimal import Decimal
import os
from pathlib import Path
from threading import Event
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from uuid import uuid4
from urllib.parse import unquote, urlsplit

from docx import Document as WordDocument
from fastapi import HTTPException, FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select, text, update
from sqlalchemy.engine import make_url
from sqlalchemy.orm import sessionmaker

from app.core.departments.models import Department
from app.core.users.models import User
from app.core.deps import get_current_user, get_db
from app.core.database import Base
from app.core.pdf.converter import PDFConversionError
from app.modules.esolver_export.service import list_esolver_pdf_certificates
from app.modules.quarta_taglio import service
from app.modules.quarta_taglio.ddt_snapshot import _normalize
from app.modules.quarta_taglio.ddt_context import resolve_saved_ddt, certificate_for_saved_ddt
from app.modules.quarta_taglio.models import QuartaTaglioDdtWorkItem, QuartaTaglioFinalCertificate, QuartaTaglioEsolverLink, QuartaTaglioCertificatePdfVersion
from app.modules.quarta_taglio.schemas import QuartaTaglioDetailResponse
from app.modules.quarta_taglio.router import router
from test_quarta_taglio_ddt_queue import DdtQueueFixture


class DdtHistoryTest(DdtQueueFixture):
    def setUp(self):
        super().setUp()
        department = Department(name="Qualità", description="Test")
        self.db.add(department)
        self.db.flush()
        self.actor = User(name="Test", email="test@example.invalid", role="manager", department_id=department.id)
        self.db.add(self.actor)
        self.db.commit()
        # No company connection is ever allowed in these tests.
        for name in ("_fetch_esolver_ddt_rows_batch",):
            guard = patch.object(service, name, side_effect=AssertionError("Unexpected remote read"))
            guard.start()
            self.addCleanup(guard.stop)
        certiol_guard = patch.object(service, "_fetch_certiol_rows_batch", return_value={})
        self.certiol = certiol_guard.start()
        self.addCleanup(certiol_guard.stop)

    def cert(self, *args, **kwargs):
        cert = super().cert(*args, **kwargs)
        if cert.storage_key_docx:
            WordDocument().save(Path(self.files.name, cert.storage_key_docx))
        return cert

    def detail(self, item, **kwargs):
        return service.get_quarta_taglio_detail(self.db, cod_odp=item.cod_odp,
                                               ddt_work_item_id=item.id, **kwargs)

    @contextmanager
    def allow_word(self):
        # Gate logic has its own tests; this fixture makes only its outcome positive.
        def build(**kwargs):
            path = kwargs["output_path"]
            path.parent.mkdir(parents=True, exist_ok=True)
            doc = WordDocument()
            doc.add_paragraph(kwargs["detail"].header["ddt"] or "no DDT")
            doc.save(path)
        with ExitStack() as stack:
            stack.enter_context(patch.object(service, "_word_creation_blockers", return_value=[]))
            stack.enter_context(patch.object(service, "_build_standard_conformity", return_value=("conforme", [])))
            stack.enter_context(patch.object(service, "build_forgialluminio_draft_docx", side_effect=build))
            yield

    def create(self, item, **kwargs):
        return service.create_quarta_taglio_word_draft(self.db, cod_odp=item.cod_odp, actor=self.actor,
                                                      ddt_work_item_id=item.id, **kwargs)

    def test_old_ddt_opens_its_own_quota_with_no_remote_reads_or_writes(self):
        item = self.item(DDT="1-01/01/2026")
        item.source_present = False
        self.material()
        self.item(IdDocumento="101", DDT="2-29/09/2026")
        self.db.add(QuartaTaglioEsolverLink(cod_odp="OL1", rows=[{"ddt": "LIVE-DDT"}], status="ok"))
        self.db.commit()
        operations = []
        def capture(conn, cursor, sql, params, context, many):
            operations.append(sql.strip().split()[0].upper())
        event.listen(self.engine, "before_cursor_execute", capture)
        try:
            detail = self.detail(item)
        finally:
            event.remove(self.engine, "before_cursor_execute", capture)
        self.assertEqual(detail.header["ddt"], item.ddt_raw)
        self.assertEqual(detail.header["unit_key"], item.certification_unit_key)
        self.assertEqual(detail.header["quantita"], "42")
        self.assertFalse(detail.ddt_source_present)
        self.assertEqual(len(detail.certifiable_units), 1)
        self.assertTrue(all(operation == "SELECT" for operation in operations))
        self.assertFalse(self.db.new or self.db.dirty)
        self.assertEqual(self.db.scalar(select(QuartaTaglioEsolverLink)).rows, [{"ddt": "LIVE-DDT"}])

    def test_finished_f3_is_not_replaced_by_raw_candidate(self):
        item = self.item(CodF3="001260")
        self.material()
        detail = self.detail(item)
        self.assertEqual(detail.header["codice_f3"], "001260")
        self.assertEqual(detail.header["ddt_finished"], item.ddt_raw)
        self.assertEqual(detail.header["codice_f3_raw"], "001230")

    def test_certiol_still_enriches_finished_description_without_ddt_query(self):
        item = self.item(CodF3="001260")
        self.material()
        self.certiol.return_value = {"OL1": [service._CertiOlRow(orp="OL1", cod_cli=None, rag_soc="Cliente A",
                                                              cod_f3_odp="001230", cod_f3="001260", des_f3="Particolare finito") ]}
        detail = self.detail(item)
        self.assertEqual(detail.header["descrizione_finished"], "Particolare finito")
        self.assertEqual(detail.header["ddt_finished"], item.ddt_raw)

    def test_ddt_number_not_misread_as_day_and_date_formats_preserved(self):
        for text in ("77-01/09/2026", "12-01/09/2026", "1-01/09/2026", "1234-01/09/2026",
                     "01/09/2026", "01-09-2026", "12-01-09-2026", "01/09/26", "01-09-26"):
            with self.subTest(text=text):
                self.assertEqual(service._certificate_datetime_from_ddt(text), datetime(2026, 9, 1, tzinfo=timezone.utc))
        self.assertIsNone(service._certificate_datetime_from_ddt("31/02/2026"))
        self.assertIsNone(service._certificate_datetime_from_ddt("nessuna data"))

    def test_no_ol_wrong_ol_or_missing_quarta_are_not_invented(self):
        item = self.item()
        with self.assertRaises(HTTPException) as exc:
            service.get_quarta_taglio_detail(self.db, cod_odp="OTHER", ddt_work_item_id=item.id)
        self.assertEqual(exc.exception.status_code, 404)
        with self.assertRaises(HTTPException) as exc:
            self.detail(item)
        self.assertEqual(exc.exception.status_code, 404)

    def test_ambiguous_identity_or_missing_date_blocks_saved_selection(self):
        item = self.item()
        self.material()
        for reason in ("source_identity_changed", "ddt_date_unrecognized"):
            item.source_review_reason = reason
            self.db.commit()
            with self.assertRaises(HTTPException) as exc:
                self.detail(item)
            self.assertEqual(exc.exception.status_code, 409)

    def test_conflicting_candidate_parameter_rejected(self):
        item = self.item()
        with self.assertRaises(HTTPException) as exc:
            self.detail(item, candidate_cod_f3="other")
        self.assertEqual(exc.exception.status_code, 422)
        with self.assertRaises(HTTPException) as exc:
            self.create(item, candidate_cod_f3="other")
        self.assertEqual(exc.exception.status_code, 422)

    def test_legacy_same_ddt_certificate_does_not_get_assigned(self):
        item = self.item()
        self.material()
        cert = self.cert(item, unit_key=None, esolver_id_documento=None, esolver_id_riga_doc=None)
        with self.assertRaises(HTTPException):
            self.detail(item)
        self.assertIsNone(cert.unit_key)

    def test_early_word_preview_has_selected_ddt_but_does_not_reassign_word(self):
        item = self.item()
        self.material()
        cert = self.cert(item, ddt=None, unit_key="early", esolver_id_documento=None,
                         esolver_id_riga_doc=None, esolver_rif_lotto_alfanum=None)
        detail = self.detail(item)
        self.assertEqual(detail.ddt_early_word_id, cert.id)
        self.assertEqual(detail.header["ddt"], item.ddt_raw)
        self.assertIsNone(detail.header["certificate_id"])
        self.assertEqual(cert.unit_key, "early")
        self.assertFalse(self.db.dirty)

    def test_early_word_multiple_quotas_is_a_conflict(self):
        item = self.item()
        self.material()
        self.cert(item, ddt=None, unit_key="early", esolver_id_documento=None,
                  esolver_id_riga_doc=None, esolver_rif_lotto_alfanum=None)
        self.item(IdDocumento="101")
        with self.assertRaises(HTTPException) as exc:
            self.detail(item)
        self.assertEqual(exc.exception.status_code, 409)

    def test_real_standard_gate_still_blocks_word_creation(self):
        item = self.item()
        self.material()
        detail = self.detail(item)
        self.assertFalse(detail.can_create_word)
        with self.assertRaises(HTTPException):
            self.create(item)
        self.assertEqual(list(self.db.scalars(select(QuartaTaglioFinalCertificate))), [])

    def test_empty_quarta_cdq_shows_blocker_not_crash(self):
        item = self.item()
        self.material(cdq="")
        detail = self.detail(item)
        self.assertFalse(detail.ready)
        self.assertFalse(detail.can_create_word)

    def test_missing_shipped_quantity_never_falls_back_to_material_weight(self):
        item = self.item(QtaUmMag=None)
        self.material()
        self.assertIsNone(self.detail(item).header["quantita"])
        with self.allow_word():
            result = self.create(item)
        self.assertIsNone(self.db.get(QuartaTaglioFinalCertificate, result.id).quantita)

    def test_register_uses_saved_shipment_with_no_odvf3(self):
        item = self.item(ODVF3=None)
        self.material()
        with self.allow_word():
            result = self.create(item)
        cert = self.db.get(QuartaTaglioFinalCertificate, result.id)
        registered = service._serialize_final_certificate_register_items(self.db, [cert])[0]
        self.assertEqual(registered.quantita, 42)

    def test_live_register_quantity_not_overwritten_by_older_snapshot(self):
        item = self.item()
        self.material()
        with self.allow_word():
            result = self.create(item)
        cert = self.db.get(QuartaTaglioFinalCertificate, result.id)
        self.db.add(QuartaTaglioEsolverLink(cod_odp="OL1", status="ok", rows=[{
            "id_documento": item.id_documento, "id_riga_doc": item.id_riga_doc,
            "rif_lotto_alfanum": item.rif_lotto_alfanum, "orp": item.cod_odp, "cod_f3": item.cod_f3,
            "ddt": item.ddt_raw, "odv_cli": item.ordine_cliente, "odv_f3": item.conferma_ordine, "qta_um_mag": 99,
        }]))
        self.db.commit()
        self.assertEqual(service._serialize_final_certificate_register_items(self.db, [cert])[0].quantita, 99)
        self.assertEqual(cert.quantita, 42)

    def test_word_creation_keeps_exact_identity_and_does_not_create_siblings(self):
        item = self.item()
        self.item(IdDocumento="101", IdRigaDoc="2", QtaUmMag=7)
        self.material()
        with self.allow_word():
            result = self.create(item)
        cert = self.db.get(QuartaTaglioFinalCertificate, result.id)
        self.assertEqual((cert.unit_key, cert.esolver_id_documento, cert.esolver_id_riga_doc),
                         (item.certification_unit_key, "100", "1"))
        self.assertEqual(cert.quantita, 42)
        self.assertEqual(len(list(self.db.scalars(select(QuartaTaglioFinalCertificate)))), 1)
        detail = service.get_quarta_taglio_detail(self.db, cod_odp="OL1", certificate_id=cert.id)
        self.assertEqual(detail.ddt_work_item_id, item.id)
        self.assertEqual(detail.header["ddt"], item.ddt_raw)

    def test_repeated_generation_reuses_record_and_unique_early_word(self):
        item = self.item()
        self.material()
        early = self.cert(item, ddt=None, unit_key="early", esolver_id_documento=None,
                          esolver_id_riga_doc=None, esolver_rif_lotto_alfanum=None)
        with self.allow_word():
            first = self.create(item)
            second = self.create(item)
        self.assertEqual((first.id, second.id), (early.id, early.id))
        self.assertEqual(early.unit_key, item.certification_unit_key)
        self.assertEqual(len(list(self.db.scalars(select(QuartaTaglioFinalCertificate)))), 1)

    def test_closed_pdf_cannot_be_regenerated_and_is_not_modified_on_view(self):
        item = self.item()
        self.material()
        cert = self.cert(item, pdf=True)
        before = {column.key: getattr(cert, column.key) for column in cert.__table__.columns}
        self.detail(item)
        self.assertEqual(before, {column.key: getattr(cert, column.key) for column in cert.__table__.columns})
        with self.allow_word(), self.assertRaises(HTTPException) as exc:
            self.create(item)
        self.assertEqual(exc.exception.status_code, 409)

    def test_manual_early_word_requires_existing_regeneration_confirmation(self):
        item = self.item()
        self.material()
        cert = self.cert(item, ddt=None, unit_key="early", esolver_id_documento=None,
                         esolver_id_riga_doc=None, esolver_rif_lotto_alfanum=None, word_source="user_uploaded")
        with self.allow_word(), self.assertRaises(HTTPException) as exc:
            self.create(item)
        self.assertEqual(exc.exception.status_code, 409)
        self.db.rollback()
        self.assertEqual(cert.unit_key, "early")

    def test_changed_source_blocks_download_even_manual_word_without_controls(self):
        item = self.item()
        cert = self.cert(item)
        item.quantita = Decimal("99")
        self.db.commit()
        with self.assertRaises(HTTPException) as exc:
            service._sync_word_fields_for_download(self.db, certificate=cert,
                                                   path=Path(self.files.name, cert.storage_key_docx))
        self.assertEqual(exc.exception.status_code, 409)
        self.assertEqual(cert.quantita, 42)

    def test_no_matching_snapshot_keeps_existing_certificate_path(self):
        item = self.item()
        cert = self.cert(item, esolver_id_documento="other", unit_key="old")
        self.assertIsNone(resolve_saved_ddt(self.db, cod_odp="OL1", certificate=cert))

    def test_source_correction_cannot_fall_back_to_other_ddt(self):
        item = self.item()
        cert = self.cert(item)
        item.ddt_raw = "changed"
        item.certification_unit_key = "changed"
        self.db.commit()
        with self.assertRaises(HTTPException):
            resolve_saved_ddt(self.db, cod_odp="OL1", certificate=cert)

    def test_download_with_controls_uses_copy_on_write(self):
        item = self.item()
        self.material()
        cert = self.cert(item)
        original = Path(self.files.name, cert.storage_key_docx)
        before = original.read_bytes()
        def update(source, target, values):
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(b"updated quota Word")
            self.assertEqual(values["DDT_RAW"], item.ddt_raw)
        with patch.object(service, "inspect_docx_content_controls", return_value=(["DDT_RAW"], [])), \
             patch.object(service, "update_docx_content_controls", side_effect=update):
            updated = service._sync_word_fields_for_download(self.db, certificate=cert, path=original)
        self.assertNotEqual(updated, original)
        self.assertEqual(original.read_bytes(), before)
        self.assertEqual(cert.ddt, item.ddt_raw)

    def test_pdf_closure_and_reopen_affect_only_selected_saved_quota(self):
        item = self.item()
        other = self.item(IdDocumento="101")
        self.material()
        with self.allow_word():
            result = self.create(item)
        def convert(word, pdf):
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b"%PDF-1.4 mocked external converter")
        with patch.object(service, "convert_docx_to_pdf", side_effect=convert):
            registered = service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=result.id, actor=self.actor)
        self.assertEqual(registered.quantita, 42)
        self.assertEqual([row.id for row in self.read().items], [other.id])
        service.reopen_quarta_taglio_certificate_pdf(self.db, certificate_id=result.id, reason="Test", actor=self.actor)
        self.assertEqual(self.read().active, 2)

    def test_source_change_during_conversion_prevents_pdf_closure(self):
        item = self.item()
        self.material()
        with self.allow_word():
            result = self.create(item)
        def convert(word, pdf):
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b"%PDF mock")
            item.quantita = Decimal("999")
            self.db.commit()
        with patch.object(service, "convert_docx_to_pdf", side_effect=convert), self.assertRaises(HTTPException):
            service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=result.id, actor=self.actor)
        self.assertEqual(self.db.get(QuartaTaglioFinalCertificate, result.id).status, "draft")
        self.assertIsNone(self.db.get(QuartaTaglioFinalCertificate, result.id).pdf_file_name)

    def test_named_pdf_route_persists_export_download_and_version_history(self):
        item = self.item()
        self.material()
        with self.allow_word():
            draft = self.create(item)
        certificate_number = self.db.get(QuartaTaglioFinalCertificate, draft.id).certificate_number
        app = FastAPI()
        app.include_router(router, prefix="/api/quarta-taglio")
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.actor
        client = TestClient(app)
        path = f"/api/quarta-taglio/certificates/{draft.id}"

        def convert(word, pdf):
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b"%PDF test converter output")

        with patch.object(service, "convert_docx_to_pdf", side_effect=convert) as converter:
            response = client.post(path + "/pdf", json={"pdf_file_name": "Cliente qualità - ordine 123.PDF"})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["pdf_file_name"], "Cliente qualità - ordine 123.pdf")
            self.db.expire_all()
            export = list_esolver_pdf_certificates(self.db, public_base_url="http://testserver").items[0]
            self.assertEqual(export.nome_file_pdf, "Cliente qualità - ordine 123.pdf")
            self.assertEqual(export.numero_certificato, certificate_number)
            url = urlsplit(export.pdf_url)
            download = client.get(url.path + "?" + url.query)
            self.assertEqual(download.status_code, 200)
            self.assertEqual(download.content, b"%PDF test converter output")
            self.assertIn(export.nome_file_pdf, unquote(download.headers["content-disposition"]))
            self.assertEqual(client.post(path + "/pdf").status_code, 200)  # Legacy request is idempotent.
            self.assertEqual(client.post(path + "/pdf", json={"pdf_file_name": "changed.pdf"}).status_code, 409)
            self.assertEqual(converter.call_count, 1)

            reopened = client.post(path + "/reopen", json={"reason": "Test nuova versione"})
            self.assertEqual(reopened.status_code, 200, reopened.text)
            self.assertEqual(reopened.json()["pdf_file_name"], "Cliente qualità - ordine 123.pdf")
            self.assertEqual(list_esolver_pdf_certificates(self.db, public_base_url="http://testserver").total_items, 0)
            regenerated = client.post(path + "/pdf", json={"pdf_file_name": "Nuovo nome"})
            self.assertEqual(regenerated.status_code, 200, regenerated.text)
            self.assertEqual(regenerated.json()["pdf_file_name"], "Nuovo nome.pdf")
            self.assertEqual(regenerated.json()["certificate_number"], certificate_number)

        versions = self.db.scalars(select(QuartaTaglioCertificatePdfVersion).order_by(QuartaTaglioCertificatePdfVersion.version)).all()
        self.assertEqual([(v.status, v.pdf_file_name) for v in versions], [
            ("reopened", "Cliente qualità - ordine 123.pdf"), ("active", "Nuovo nome.pdf"),
        ])
        self.assertNotEqual(versions[0].storage_key_pdf, versions[1].storage_key_pdf)
        self.assertEqual(list_esolver_pdf_certificates(self.db, public_base_url="http://testserver").items[0].nome_file_pdf, "Nuovo nome.pdf")

    def test_failed_conversion_does_not_save_new_name_or_close_certificate(self):
        item = self.item()
        self.material()
        with self.allow_word():
            draft = self.create(item)
        with patch.object(service, "convert_docx_to_pdf", side_effect=PDFConversionError("Test failure")):
            with self.assertRaises(HTTPException) as error:
                service.generate_quarta_taglio_certificate_pdf(
                    self.db, certificate_id=draft.id, actor=self.actor, pdf_file_name="Cliente.pdf",
                )
        self.assertEqual(error.exception.status_code, 502)
        self.db.expire_all()
        certificate = self.db.get(QuartaTaglioFinalCertificate, draft.id)
        self.assertIsNone(certificate.pdf_file_name)
        self.assertIsNone(certificate.storage_key_pdf)
        self.assertEqual(certificate.status, "draft")
        self.assertEqual(self.db.scalars(select(QuartaTaglioCertificatePdfVersion)).all(), [])

    def test_saved_routes_accept_positive_id_and_enforce_department(self):
        item = self.item()
        self.material()
        app = FastAPI()
        app.include_router(router, prefix="/api/quarta-taglio")
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.actor
        client = TestClient(app)
        path = "/api/quarta-taglio/OL1"
        response = client.get(path, params={"ddt_work_item_id": item.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["ddt_work_item_id"], item.id)
        self.assertEqual(client.get(path, params={"ddt_work_item_id": -1}).status_code, 422)
        self.actor.department.name = "Produzione"
        self.assertEqual(client.get(path, params={"ddt_work_item_id": item.id}).status_code, 403)
        self.assertEqual(client.post(path + "/word-draft", json={"ddt_work_item_id": item.id}).status_code, 403)


@unittest.skipUnless(os.environ.get("DDT_TEST_POSTGRES_URL"), "isolated PostgreSQL test URL not supplied")
class DdtHistoryPostgresTest(unittest.TestCase):
    def setUp(self):
        url = make_url(os.environ["DDT_TEST_POSTGRES_URL"])
        if url.host not in ("127.0.0.1", "localhost") or not (url.database or "").startswith("certi_ddt_test"):
            self.fail("Only a dedicated local certi_ddt_test database is permitted")
        self.root_engine = create_engine(url)
        self.schema = "ddt_history_test_" + uuid4().hex
        with self.root_engine.begin() as conn:
            conn.execute(text(f'CREATE SCHEMA "{self.schema}"'))
        self.engine = self.root_engine.execution_options(schema_translate_map={None: self.schema})
        Base.metadata.create_all(self.engine)
        self.factory = sessionmaker(bind=self.engine, autoflush=False)
        with self.factory.begin() as db:
            values = _normalize(dict(IdDocumento="1", IdRigaDoc="1", RifLottoAlfanum="L1", ORP="OL1", CodF3="001230",
                                     DDT="77-01/09/2026", RagSoc="Test", ODVCli="PO", ODVF3="ORD", QtaUmMag=42, CertificatoPresente=0))
            item = QuartaTaglioDdtWorkItem(**values, first_seen_at=datetime.now(timezone.utc), last_seen_at=datetime.now(timezone.utc))
            db.add(item)
            db.flush()
            self.item_id = item.id

    def tearDown(self):
        with self.root_engine.begin() as conn:
            conn.execute(text(f'DROP SCHEMA "{self.schema}" CASCADE'))
        self.root_engine.dispose()

    def create_record(self, *, acquired=None, release=None, attempted=None):
        with self.factory.begin() as db:
            if attempted:
                attempted.set()
            item = resolve_saved_ddt(db, cod_odp="OL1", work_item_id=self.item_id, lock=True)
            service._lock_certificate_register_for_ol(db, cod_odp="OL1")
            if acquired:
                acquired.set()
                if not release.wait(5):
                    raise AssertionError("Test lock release timed out")
            from app.modules.quarta_taglio.ddt_context import saved_ddt_row
            unit = service._build_certifiable_units(cod_odp="OL1", esolver_rows=[saved_ddt_row(item)], quarta_rows=[])[0]
            detail = QuartaTaglioDetailResponse(cod_odp="OL1", ready=True, status_color="green", status_message="Test",
                       header={"unit_key": unit.unit_key, "ddt": unit.ddt}, materials=[], missing_items=[],
                       standard_candidates=[], chemistry=[], properties=[], notes=[], certifiable_units=[unit])
            cert = service._get_or_create_open_certificate_for_unit(db, detail=detail, unit=unit,
                        actor=SimpleNamespace(id=None), saved_ddt=item)
            return cert.id

    def test_two_concurrent_creators_reuse_one_record(self):
        acquired, release, attempted = Event(), Event(), Event()
        with ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(self.create_record, acquired=acquired, release=release)
            try:
                self.assertTrue(acquired.wait(5))
                second = pool.submit(self.create_record, attempted=attempted)
                self.assertTrue(attempted.wait(5))
                with self.assertRaises(TimeoutError):
                    second.result(timeout=0.2)
            finally:
                release.set()
            self.assertEqual(first.result(timeout=5), second.result(timeout=5))
        with self.factory() as db:
            self.assertEqual(len(list(db.scalars(select(QuartaTaglioFinalCertificate)))), 1)

    def test_snapshot_update_waits_then_changed_quantity_is_detected(self):
        acquired, release, attempted = Event(), Event(), Event()
        def update_source():
            with self.factory.begin() as db:
                attempted.set()
                db.execute(update(QuartaTaglioDdtWorkItem).where(QuartaTaglioDdtWorkItem.id == self.item_id)
                           .values(quantita=99))
        with ThreadPoolExecutor(max_workers=2) as pool:
            creator = pool.submit(self.create_record, acquired=acquired, release=release)
            try:
                self.assertTrue(acquired.wait(5))
                updater = pool.submit(update_source)
                self.assertTrue(attempted.wait(5))
                with self.assertRaises(TimeoutError):
                    updater.result(timeout=0.2)
            finally:
                release.set()
            cert_id = creator.result(timeout=5)
            updater.result(timeout=5)
        with self.factory() as db:
            cert = db.get(QuartaTaglioFinalCertificate, cert_id)
            self.assertEqual(cert.quantita, 42)
            with self.assertRaises(HTTPException) as exc:
                resolve_saved_ddt(db, cod_odp="OL1", certificate=cert)
            self.assertEqual(exc.exception.status_code, 409)
