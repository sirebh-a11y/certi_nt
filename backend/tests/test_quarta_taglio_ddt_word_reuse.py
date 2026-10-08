"""Local-only regression coverage: Word once, one immutable PDF per received quota."""
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch
import os
import unittest
from tempfile import TemporaryDirectory
from concurrent.futures import ThreadPoolExecutor

from sqlalchemy import select, event

from app.modules.quarta_taglio import service
from app.modules.quarta_taglio import ddt_word_reuse as reuse
from app.modules.quarta_taglio.ddt_word_recovery import build_plan, apply_plan, validate_plan
from app.modules.quarta_taglio.models import QuartaTaglioFinalCertificate as Certificate, QuartaTaglioCertificatePdfAttachment
from app.modules.quarta_taglio.schemas import QuartaTaglioCodF3CandidateResponse
import test_quarta_taglio_ddt_history as history


class DdtWordReuseTest(history.DdtHistoryTest):
    # Inherit fixture helpers, not the parent test suite.
    def base_word(self, *, closed=True):
        first = self.item(DDT='2354-01/07/2026', QtaUmMag=3000)
        self.material()
        detail = self.detail(first)
        source = self.cert(first, pdf=closed, certificate_number='7013_00_00/26',
                           draft_number='7013_00_00/26', word_source='generated',
                           cdq_key=service._cdq_key_from_detail(detail),
                           cdq_signature=service._cdq_signature_from_detail(detail),
                           cdq_values=service._cdq_values_from_detail(detail))
        path = service._certificate_storage_path(source.storage_key_docx)
        service.build_forgialluminio_draft_docx(detail=detail, output_path=path,
            draft_number=source.certificate_number, certified_by=self.actor, quality_manager=self.actor)
        self.db.commit()
        return first, source

    def later(self, **kwargs):
        values = dict(IdDocumento='200', IdRigaDoc='2', DDT='2397-05/10/2026', QtaUmMag=1398,
                      ODVCli='NEW-ORDER', ODVF3='NEW-CONFIRM')
        values.update(kwargs)
        return self.item(**values)

    def run_worker(self):
        result = reuse.sync_ddt_words(session_factory=self.factory)
        self.db.expire_all()
        return result

    def test_late_ddt_after_closed_pdf_reuses_word_and_preserves_source(self):
        first, source = self.base_word()
        old_word = reuse.file_hash(source.storage_key_docx)
        old_pdf = reuse.file_hash(source.storage_key_pdf)
        second = self.later()
        self.assertEqual(self.run_worker()['prepared'], 1)
        target = self.db.scalar(select(Certificate).where(Certificate.unit_key == second.certification_unit_key))
        self.assertEqual(target.certificate_number, source.certificate_number)
        self.assertEqual(target.word_source, 'ddt_reused')
        self.assertEqual(target.quantita, 1398)
        self.assertIsNone(target.storage_key_pdf)
        controls, _ = reuse.word_facts(target.storage_key_docx)
        self.assertEqual(controls['DDT_RAW'], second.ddt_raw)
        self.assertEqual(controls['QUANTITY_RAW'], '1398')
        self.assertEqual(controls['ORDER_CLIENT'], 'NEW-ORDER')
        self.assertEqual(controls['CERT_DATE'], '05/10/2026')
        self.assertEqual(reuse.file_hash(source.storage_key_docx), old_word)
        self.assertEqual(reuse.file_hash(source.storage_key_pdf), old_pdf)
        token = target.download_token
        self.assertEqual(self.run_worker()['prepared'], 0)
        self.assertEqual(target.download_token, token)
        self.assertEqual(len(service.list_quarta_taglio_final_certificates(self.db).items), 2)
        self.assertEqual(service._sync_word_fields_for_download(self.db, certificate=target,
            path=service._certificate_storage_path(target.storage_key_docx)), service._certificate_storage_path(target.storage_key_docx))

    def test_existing_empty_record_is_filled_not_duplicated(self):
        _, source = self.base_word()
        second = self.later()
        empty = self.cert(second, word=False, certificate_number=source.certificate_number)
        self.assertEqual(self.run_worker()['prepared'], 1)
        self.assertIsNotNone(empty.storage_key_docx)
        self.assertEqual(self.db.query(Certificate).count(), 2)

    def test_no_temporal_cutoff_for_historical_quota(self):
        self.base_word()
        second = self.later(DDT='22-02/07/2026')
        second.source_present = False
        self.db.commit()
        self.assertEqual(self.run_worker()['prepared'], 1)

    def test_early_base_and_two_deliveries_remain_separate(self):
        first, source = self.base_word(closed=False)
        source.ddt = source.esolver_id_documento = source.esolver_id_riga_doc = source.esolver_rif_lotto_alfanum = None
        source.unit_key = 'early'
        self.db.commit()
        self.later()
        self.assertEqual(self.run_worker()['prepared'], 2)
        self.assertEqual(self.db.query(Certificate).count(), 3)
        self.assertIsNone(source.ddt)

    def test_missing_ol_is_not_inferred_from_numeric_lot_or_article(self):
        self.base_word()
        second = self.later(ORP=None, RifLottoAlfanum='2026000997')
        self.assertEqual(reuse.plan_item(self.db, second)['reason'], 'missing_ol_from_esolver')
        self.assertEqual(self.run_worker()['prepared'], 0)

    def test_different_workmanship_is_not_automatically_selected(self):
        self.base_word()
        second = self.later(CodF3='001260')
        self.assertEqual(reuse.plan_item(self.db, second)['reason'], 'word_not_prepared')

    def test_changed_heat_prevents_automatic_reuse(self):
        _, source = self.base_word()
        source.cdq_values = [dict(v, colata='DIFFERENT') for v in source.cdq_values]
        self.db.commit()
        second = self.later()
        self.assertEqual(reuse.plan_item(self.db, second)['reason'], 'no_compatible_word')

    def test_changed_closed_quantity_is_never_repaired(self):
        first, source = self.base_word()
        first.quantita = 111
        self.db.commit()
        self.assertEqual(reuse.plan_item(self.db, first)['reason'], 'quantity_changed')

    def test_missing_controls_are_review_not_fresh_generation(self):
        _, source = self.base_word()
        from docx import Document
        Document().save(service._certificate_storage_path(source.storage_key_docx))
        second = self.later()
        self.assertEqual(reuse.plan_item(self.db, second)['reason'], 'source_file_or_controls_invalid')

    def test_file_failure_rolls_back_new_record(self):
        self.base_word()
        self.later()
        with patch.object(reuse, 'update_docx_content_controls', side_effect=OSError('test')):
            self.assertEqual(self.run_worker()['errors'], 1)
        self.assertEqual(self.db.query(Certificate).count(), 1)

    def test_preview_is_read_only_and_report_detects_changed_quantity(self):
        self.base_word()
        second = self.later()
        target = dict(test='local')
        statements = []
        def capture(conn, cursor, sql, params, context, many):
            statements.append(sql.strip().split()[0])
        event.listen(self.engine, 'before_cursor_execute', capture)
        try:
            report = build_plan(self.db, target=target)
        finally:
            event.remove(self.engine, 'before_cursor_execute', capture)
        self.assertTrue(all(s == 'SELECT' for s in statements))
        second.quantita = 999
        self.db.commit()
        with self.assertRaises(ValueError):
            validate_plan(report, build_plan(self.db, target=target))

    def test_recovery_is_repeatable_and_preserves_closed_files(self):
        _, source = self.base_word()
        self.later()
        target = dict(test='local')
        report = build_plan(self.db, target=target)
        paths = []
        changes = apply_plan(self.db, approved=report, target=target, created_paths=paths)
        self.db.commit()
        self.assertEqual(len(changes), 1)
        self.assertTrue(paths[0].is_file())
        self.assertNotIn('reuse', build_plan(self.db, target=target)['summary'])

    def test_progress_describes_received_ddts_not_final_ol_closure(self):
        first, source = self.base_word()
        rows = service._load_quarta_rows_for_detail(self.db, cod_odp=first.cod_odp)
        kwargs = dict(group_rows=rows, certiol_rows=[], esolver_link=None, certificates=[source])
        progress = service._certification_progress_for_group(**kwargs, saved_items=[first])
        self.assertEqual(progress.label, 'PDF pronti per tutti i DDT ricevuti')
        second = self.later()
        progress = service._certification_progress_for_group(**kwargs, saved_items=[first, second])
        self.assertEqual(progress.label, 'PDF da completare per DDT ricevuti')

    def test_generic_page_finds_closed_word_despite_newer_empty_record(self):
        first, source = self.base_word()
        second = self.later()
        self.cert(second, word=False, certificate_number=source.certificate_number)
        from app.modules.quarta_taglio.ddt_context import saved_ddt_row
        link = history.QuartaTaglioEsolverLink(cod_odp='OL1', status='ok',
            rows=[saved_ddt_row(i).model_dump() for i in (first, second)])
        self.db.add(link)
        self.db.commit()
        with patch.object(service, '_refresh_quarta_rows_from_incoming'), patch.object(
            service, '_refresh_esolver_links_for_rows', return_value={'OL1':link}):
            detail = service.get_quarta_taglio_detail(self.db, cod_odp='OL1')
        self.assertEqual(detail.header['certificate_id'], str(source.id))
        self.assertTrue(detail.word_info.has_word)
        self.assertTrue(detail.word_info.is_pdf_final)
        self.assertEqual(detail.status_message, 'PDF pronto per questo DDT')
        self.assertEqual(self.db.query(Certificate).count(), 2)

    def test_manual_text_and_attachments_are_preserved(self):
        _, source = self.base_word(closed=False)
        from docx import Document
        path = service._certificate_storage_path(source.storage_key_docx)
        document = Document(path)
        document.add_paragraph('TESTO MANUALE DA CONSERVARE')
        document.save(path)
        source.word_source = 'user_uploaded'
        source.pdf_attachments_initialized = True
        Path(self.files.name, 'attachment.pdf').write_bytes(b'%PDF-1.4 isolated test')
        self.db.add(QuartaTaglioCertificatePdfAttachment(certificate_id=source.id,
            certificate_number=source.certificate_number, cod_odp=source.cod_odp,
            storage_key_pdf='attachment.pdf',original_filename='attachment.pdf',sort_order=0))
        self.db.commit()
        later = self.later()
        self.assertEqual(self.run_worker()['prepared'], 1)
        target = self.db.scalar(select(Certificate).where(Certificate.unit_key == later.certification_unit_key))
        self.assertIn('TESTO MANUALE DA CONSERVARE', '\n'.join(p.text for p in Document(service._certificate_storage_path(target.storage_key_docx)).paragraphs))
        attachments = service._pdf_attachments_for_certificate(self.db, certificate=target)
        self.assertEqual([a.storage_key_pdf for a in attachments], ['attachment.pdf'])

    def test_conflicting_manual_sources_are_not_chosen_by_recency(self):
        first, source = self.base_word()
        second = self.later()
        other = self.cert(second, certificate_number=source.certificate_number,
            cdq_values=source.cdq_values, draft_number=source.certificate_number)
        import shutil
        from docx import Document
        path = service._certificate_storage_path(other.storage_key_docx)
        shutil.copyfile(service._certificate_storage_path(source.storage_key_docx), path)
        doc = Document(path)
        doc.add_paragraph('DIVERSA NOTA TECNICA')
        doc.save(path)
        third = self.item(IdDocumento='300', DDT='300-06/10/2026')
        self.assertEqual(reuse.plan_item(self.db, third)['reason'], 'different_word_sources')

    def test_incoming_pending_is_not_automatically_confirmed(self):
        _, source = self.base_word()
        from app.modules.acquisition.models import AcquisitionRow
        incoming = self.db.scalar(select(AcquisitionRow))
        incoming.qualita_valutazione = None
        self.db.commit()
        later = self.later()
        self.assertEqual(reuse.plan_item(self.db, later)['reason'], 'incoming_not_ready')
        self.assertEqual(self.run_worker()['prepared'], 0)
        self.assertIsNone(incoming.qualita_valutazione)

    def test_reused_word_can_close_pdf_and_export_only_its_quota(self):
        _, source = self.base_word()
        self.later()
        self.run_worker()
        target = self.db.scalar(select(Certificate).where(Certificate.id != source.id))
        target.conformity_status = 'conforme'
        self.db.commit()
        def convert(src, dest):
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_bytes(b'%PDF-1.4 local converter fixture')
        with patch.object(service, 'convert_docx_to_pdf', side_effect=convert), patch.object(
            service, '_build_standard_conformity', return_value=('conforme', [])):
            result = service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=target.id, actor=self.actor)
        self.assertTrue(result.has_pdf)
        self.assertEqual(target.ddt, '2397-05/10/2026')
        self.assertEqual(target.quantita, 1398)


# unittest also discovers inherited methods. Keep the dedicated fixture helpers,
# but don't count the existing DdtHistoryTest suite a second time here.
for _name in dir(history.DdtHistoryTest):
    if _name.startswith('test_') and _name not in DdtWordReuseTest.__dict__:
        setattr(DdtWordReuseTest, _name, None)


@unittest.skipUnless(os.environ.get('DDT_TEST_POSTGRES_URL'), 'isolated PostgreSQL test URL not supplied')
class DdtWordReusePostgresTest(DdtWordReuseTest):
    def setUp(self):
        history.DdtHistoryPostgresTest.setUp(self)
        self.addCleanup(lambda: history.DdtHistoryPostgresTest.tearDown(self))
        self.db = self.factory()
        self.addCleanup(self.db.close)
        self.db.query(history.QuartaTaglioDdtWorkItem).delete()
        self.files = TemporaryDirectory(prefix='certi-word-reuse-pg-')
        self.addCleanup(self.files.cleanup)
        storage = patch.object(service.settings, 'document_storage_root', self.files.name)
        storage.start()
        self.addCleanup(storage.stop)
        from app.modules.acquisition.models import Document
        from app.core.departments.models import Department
        from app.core.users.models import User
        department = Department(name='Qualità', description='Test isolato')
        self.db.add(department)
        self.db.flush()
        self.actor = User(name='Test', email='word@example.invalid', role='manager', department_id=department.id)
        self.document = Document(tipo_documento='certificato', nome_file_originale='test.pdf', storage_key='supplier.pdf')
        self.db.add_all([self.actor, self.document])
        self.db.commit()
        for name in ('_fetch_esolver_ddt_rows_batch', '_fetch_certiol_rows_batch'):
            guard = patch.object(service, name, return_value={})
            guard.start()
            self.addCleanup(guard.stop)

    def test_concurrent_workers_prepare_only_one_word(self):
        self.base_word()
        self.later()
        with ThreadPoolExecutor(max_workers=2) as pool:
            outcomes = list(pool.map(lambda _: reuse.sync_ddt_words(session_factory=self.factory), range(2)))
        self.assertEqual(sum(r['prepared'] for r in outcomes), 1)
        self.assertEqual(sum(r['errors'] for r in outcomes), 0)
        self.db.expire_all()
        self.assertEqual(self.db.query(Certificate).count(), 2)

    def test_offline_recovery_lock_refuses_an_active_writer(self):
        from app.modules.quarta_taglio.ddt_word_recovery import lock_inputs
        from sqlalchemy import update
        from sqlalchemy.exc import DBAPIError
        _, source = self.base_word()
        with self.factory() as writer, writer.begin():
            writer.execute(update(Certificate).where(Certificate.id == source.id).values(draft_number='LOCK-TEST'))
            with self.factory() as repair, repair.begin():
                with self.assertRaises(DBAPIError):
                    lock_inputs(repair)
            writer.rollback()

    def test_offline_recovery_lock_allows_reviewed_application(self):
        from app.modules.quarta_taglio.ddt_word_recovery import lock_inputs
        self.base_word()
        self.later()
        self.db.rollback()
        with self.factory() as repair, repair.begin():
            lock_inputs(repair)
            report = build_plan(repair, target={'test':'isolated'})
            paths = []
            changes = apply_plan(repair, approved=report, target={'test':'isolated'}, created_paths=paths)
        self.assertEqual(len(changes), 1)
