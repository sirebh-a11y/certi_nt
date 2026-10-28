"""Full lifecycle regressions: isolated DB/files, never company connections."""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import hashlib
import os
import unittest
import zipfile

from docx import Document
from docx.shared import Inches
from fastapi import HTTPException, UploadFile

from app.modules.quarta_taglio import service, word_reference as refs, ddt_word_reuse as reuse
from app.modules.quarta_taglio.certificate_docx import prepare_manual_docx_base
from app.modules.quarta_taglio.models import QuartaTaglioCertificatePdfVersion as Version
from app.modules.quarta_taglio.models import QuartaTaglioFinalCertificate as Certificate
import test_quarta_taglio_word_reference as reference_tests


class WordLifecycleTest(reference_tests.WordReferenceTest):
    def attach_pdf(self, certificate, *, actor=None):
        from pypdf import PdfWriter
        pdf = PdfWriter()
        pdf.add_blank_page(width=595, height=842)
        stream = BytesIO()
        pdf.write(stream)
        stream.seek(0)
        return service.upload_quarta_taglio_pdf_attachment(self.db,
            cod_odp=certificate.cod_odp, certificate_id=certificate.id,
            actor=actor or self.actor,
            uploaded_file=UploadFile(filename='isolated_attachment.pdf', file=stream))

    def upload(self, certificate, *, text=None, letter=False):
        doc = Document(service._certificate_storage_path(certificate.storage_key_docx))
        if text:
            doc.add_paragraph(text)
        if letter:
            doc.sections[0].page_width = Inches(8.5)
        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        return service.upload_quarta_taglio_word_file(
            self.db, cod_odp=certificate.cod_odp, certificate_id=certificate.id,
            actor=self.actor, uploaded_file=UploadFile(filename='completed.docx', file=stream))

    def test_real_upload_download_and_later_ddt_keep_valid_reference(self):
        _, source = self.base_word(closed=False)
        self.upload(source, text='MANUAL SECOND PAGE', letter=True)
        revision = refs.metadata(source)['revision']
        original_key = source.storage_key_docx
        original_hash = reuse.file_hash(original_key)
        after_first = None
        for _ in range(2):
            path = service._certificate_storage_path(source.storage_key_docx)
            service._sync_word_fields_for_download(self.db, certificate=source, path=path)
            self.assertIsNone(refs.choose(self.db, [source])[1])
            self.assertEqual(refs.metadata(source)['revision'], revision)
            self.assertEqual(refs.metadata(source)['storage_key'], source.storage_key_docx)
            if after_first is not None:
                self.assertEqual(source.storage_key_docx, after_first)
            after_first = source.storage_key_docx
        self.assertEqual(reuse.file_hash(original_key), original_hash)
        item = self.later()
        self.assertEqual(reuse.plan_item(self.db, item)['action'], 'reuse')
        self.assertEqual(self.run_worker()['prepared'], 1)
        target = self.target(item)
        self.assertEqual(refs.metadata(target)['source_id'], source.id)
        self.assertIn('MANUAL SECOND PAGE', '\n'.join(p.text for p in Document(service._certificate_storage_path(target.storage_key_docx)).paragraphs))
        # A maintenance-only rewrite of a SEPARATE source's shipment headers
        # cannot invalidate this copy's conversion state or elect a new source.
        from app.modules.quarta_taglio.certificate_docx import update_docx_content_controls
        conversion = refs.conversion_state(self.db, target)
        before = refs.before_automatic(source)
        next_key = service._certificate_docx_storage_key(source.cod_odp)
        update_docx_content_controls(service._certificate_storage_path(source.storage_key_docx),
            service._certificate_storage_path(next_key), {'DDT_RAW':'SOURCE HEADER ONLY'})
        source.storage_key_docx = next_key
        refs.after_automatic(source, before, origin='source_header')
        self.db.commit()
        self.assertEqual(refs.conversion_state(self.db, target), conversion)

    def test_multiple_ddt_preview_shows_source_without_binding_it(self):
        first, source = self.base_word(closed=False)
        source.ddt = source.esolver_id_documento = source.esolver_id_riga_doc = source.esolver_rif_lotto_alfanum = None
        source.unit_key = 'early-source'
        source.download_token = 'isolated-source-download'
        self.publish(source)
        self.later()
        before = (source.id, source.unit_key, source.storage_key_docx)
        detail = self.detail(first)
        self.assertIsNone(detail.header['certificate_id'])
        self.assertFalse(detail.word_info.has_word)
        self.assertEqual(detail.word_info.reference_certificate_id, source.id)
        self.assertTrue(detail.word_info.reference_download_url)
        self.assertEqual(detail.word_info.reference_number, source.certificate_number)
        self.assertFalse(detail.can_create_word)
        self.assertEqual((source.id, source.unit_key, source.storage_key_docx), before)
        self.assertFalse(self.db.dirty or self.db.new or self.db.deleted)

    def test_multiple_ddt_with_missing_incoming_still_opens_diagnostics(self):
        first, source = self.base_word(closed=False)
        source.ddt = source.esolver_id_documento = source.esolver_id_riga_doc = source.esolver_rif_lotto_alfanum = None
        source.unit_key = 'early-source'
        self.publish(source)
        self.material(cdq='MISSING', colata='NEW', certificate=False)
        self.later()
        detail = self.detail(first)
        self.assertFalse(detail.ready)
        self.assertTrue(detail.missing_items)
        self.assertTrue(detail.word_info.material_changed)
        self.assertIn('MISSING', detail.word_info.reference_message)
        self.assertIsNone(source.ddt)

    def test_pdf_rejects_same_reference_changed_during_conversion(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        source.conformity_status = 'conforme'
        source.cert_date = datetime.now(timezone.utc)
        self.db.commit()
        converted = []
        def convert(word, pdf):
            self.assertNotEqual(word, service._certificate_storage_path(source.storage_key_docx))
            converted.append(pdf)
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b'%PDF-1.4 audit')
            self.edit(source, 'NEW WORD DURING CONVERSION')
        with patch.object(service, 'convert_docx_to_pdf', side_effect=convert), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            with self.assertRaises(HTTPException) as caught:
                service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=source.id, actor=self.actor)
        self.assertEqual(caught.exception.status_code, 409)
        self.assertNotEqual(source.status, 'pdf_final')
        self.assertIsNone(source.storage_key_pdf)
        self.assertEqual(self.db.query(Version).count(), 0)
        self.assertTrue(converted)
        self.assertFalse(converted[0].exists())

    def test_pdf_rejects_unpublished_integrity_change_without_accepting_it(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        source.conformity_status = 'conforme'
        source.cert_date = datetime.now(timezone.utc)
        self.db.commit()
        metadata = dict(refs.metadata(source))
        path = service._certificate_storage_path(source.storage_key_docx)
        doc = Document(path)
        doc.add_paragraph('UNPUBLISHED TECHNICAL CHANGE')
        doc.save(path)
        with patch.object(service, 'convert_docx_to_pdf') as convert, patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            with self.assertRaises(HTTPException):
                service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=source.id, actor=self.actor)
        convert.assert_not_called()
        self.assertEqual(refs.metadata(source), metadata)

    def test_pdf_rejects_upload_in_gap_after_sync_before_conversion_pin(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        source.conformity_status = 'conforme'
        source.cert_date = datetime.now(timezone.utc)
        self.db.commit()
        sync = service._sync_word_fields_for_download
        def sync_then_upload(db, *, certificate, path):
            old = sync(db, certificate=certificate, path=path)
            self.upload(source, text='NEW UPLOAD AFTER SYNC')
            return old
        with patch.object(service, '_sync_word_fields_for_download', side_effect=sync_then_upload), patch.object(service, 'convert_docx_to_pdf') as convert, patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            with self.assertRaises(HTTPException) as caught:
                service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=source.id, actor=self.actor)
        self.assertEqual(caught.exception.status_code, 409)
        convert.assert_not_called()
        self.assertNotEqual(source.status, 'pdf_final')
        self.assertIsNone(source.storage_key_pdf)
        self.assertEqual(self.db.query(Version).count(), 0)

    def test_inherited_complete_word_survives_attachment_add_remove_and_later_ddt(self):
        from PIL import Image
        from app.core.users.models import User
        _, raw = self.base_word(closed=False)
        self.edit(raw, 'RAW MANUAL LAB PAGE')
        picture = Path(self.files.name) / 'lab.png'
        Image.new('RGB', (120, 50), 'green').save(picture)
        original = service._certificate_storage_path(raw.storage_key_docx)
        document = Document(original)
        document.add_picture(str(picture), width=Inches(1))
        document.save(original)
        raw.certified_by_user_id = raw.quality_manager_user_id = self.actor.id
        self.publish(raw)
        raw_hash = reuse.file_hash(raw.storage_key_docx)
        editor = User(name='Attachment editor', email='editor@example.invalid', role='manager', department_id=self.actor.department_id)
        self.db.add(editor)
        self.db.commit()
        item = self.later(CodF3='001260')
        def body(c):
            return '\n'.join(p.text for p in Document(service._certificate_storage_path(c.storage_key_docx)).paragraphs)
        def media(c):
            with zipfile.ZipFile(service._certificate_storage_path(c.storage_key_docx)) as archive:
                return {hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
        with patch.object(service, '_word_creation_blockers', return_value=[]), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            response = self.create(item)
            target = self.db.get(Certificate, response.id)
            images = media(target)
            sections_before = len(Document(service._certificate_storage_path(target.storage_key_docx)).sections)
            self.assertIn('RAW MANUAL LAB PAGE', body(target))
            self.attach_pdf(target, actor=editor)
            self.assertIn('RAW MANUAL LAB PAGE', body(target))
            self.assertTrue(images <= media(target))
            self.assertEqual((target.certified_by_user_id, target.quality_manager_user_id), (self.actor.id, self.actor.id))
            self.attach_pdf(target, actor=editor)
            self.assertEqual(len(service._pdf_attachments_for_certificate(self.db, certificate=target)), 2)
            self.assertEqual(body(target).count('CERTI_PDF_ATTACHMENT_PAGE'), 2)
            attachment = service._pdf_attachments_for_certificate(self.db, certificate=target)[0]
            service.delete_quarta_taglio_pdf_attachment(self.db, cod_odp=target.cod_odp,
                certificate_id=target.id, attachment_id=attachment.id, actor=editor)
            self.assertEqual(body(target).count('CERTI_PDF_ATTACHMENT_PAGE'), 1)
            remaining = service._pdf_attachments_for_certificate(self.db, certificate=target)[0]
            service.delete_quarta_taglio_pdf_attachment(self.db, cod_odp=target.cod_odp,
                certificate_id=target.id, attachment_id=remaining.id, actor=editor)
            self.assertNotIn('CERTI_PDF_ATTACHMENT_PAGE', body(target))
            self.assertIn('RAW MANUAL LAB PAGE', body(target))
            self.assertTrue(images <= media(target))
            service._sync_word_fields_for_download(self.db, certificate=target,
                path=service._certificate_storage_path(target.storage_key_docx))
            self.assertEqual(len(Document(service._certificate_storage_path(target.storage_key_docx)).sections), sections_before)
            following = self.later(IdDocumento='300', DDT='2500-20/10/2026', CodF3='001260')
            self.run_worker()
            copy = self.target(following)
            self.assertIsNotNone(copy)
            self.assertIn('RAW MANUAL LAB PAGE', body(copy))
            self.assertTrue(images <= media(copy))
            self.assertEqual(refs.metadata(copy)['source_id'], target.id)
            self.assertEqual(reuse.file_hash(raw.storage_key_docx), raw_hash)

    def test_completed_reused_reference_updates_total_without_changing_revision(self):
        _, raw = self.base_word()
        raw_hash, pdf_hash = reuse.file_hash(raw.storage_key_docx), reuse.file_hash(raw.storage_key_pdf)
        item = self.later(QtaUmMag=91)
        self.run_worker()
        target = self.target(item)
        with patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            self.attach_pdf(target)
            self.assertEqual(target.word_source, 'ddt_reused')
            self.assertEqual(refs.metadata(target)['kind'], 'reference')
            revision, number = refs.metadata(target)['revision'], target.certificate_number
            attachments = [a.storage_key_pdf for a in service._pdf_attachments_for_certificate(self.db, certificate=target)]
            sibling = self.later(IdRigaDoc='3', QtaUmMag=29)
            self.assertEqual(reuse.plan_item(self.db, item)['reason'], 'existing_document_preserved')
            service._sync_word_fields_for_download(self.db, certificate=target,
                path=service._certificate_storage_path(target.storage_key_docx))
            controls, _ = reuse.word_facts(target.storage_key_docx)
            self.assertEqual(controls['QUANTITY_RAW'], '91 of 120')
            self.assertEqual(refs.metadata(target)['revision'], revision)
            self.assertEqual(target.certificate_number, number)
            self.assertIsNone(refs.integrity_reason(target))
            self.assertEqual([a.storage_key_pdf for a in service._pdf_attachments_for_certificate(self.db, certificate=target)], attachments)
            self.run_worker()
            self.assertEqual(self.target(sibling).certificate_number, number)
            self.assertEqual(reuse.word_facts(self.target(sibling).storage_key_docx)[0]['QUANTITY_RAW'], '29 of 120')
            converted_controls = []
            def convert(word, pdf):
                # The published PDF must use the refreshed total, not merely
                # the record/header shown by the UI. Real LO is covered in QA.
                with zipfile.ZipFile(word) as archive:
                    converted_controls.append(''.join(archive.read(n).decode('utf-8') for n in archive.namelist() if n.startswith('word/') and n.endswith('.xml')))
                pdf.parent.mkdir(parents=True, exist_ok=True)
                pdf.write_bytes(b'%PDF-1.4 isolated refreshed quantity')
            with patch.object(service, 'convert_docx_to_pdf', side_effect=convert):
                response = service.generate_quarta_taglio_certificate_pdf(self.db,
                    certificate_id=target.id, actor=self.actor)
            self.assertTrue(response.has_pdf)
            self.assertIn('91 of 120', converted_controls[0])
            self.assertEqual(refs.metadata(target)['revision'], revision)
            self.assertEqual(reuse.file_hash(raw.storage_key_docx), raw_hash)
            self.assertEqual(reuse.file_hash(raw.storage_key_pdf), pdf_hash)

    def test_manual_bases_of_same_number_do_not_cross_documents(self):
        first, a = self.base_word(closed=False)
        second = self.later(IdDocumento=first.id_documento, DDT=first.ddt_raw)
        b = self.cert(second, certificate_number=a.certificate_number, word_source='user_uploaded')
        content = service._certificate_storage_path(a.storage_key_docx).read_bytes()
        for c, marker in [(a, 'MANUAL_CONTENT_A'), (b, 'MANUAL_CONTENT_B')]:
            key = service._certificate_docx_storage_key(c.cod_odp)
            p = service._certificate_storage_path(key)
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_bytes(content)
            doc = Document(p)
            doc.add_paragraph(marker)
            doc.save(p)
            c.storage_key_docx, c.word_source = key, 'user_uploaded'
            prepare_manual_docx_base(p, service._certificate_storage_path(service._certificate_manual_base_docx_storage_key(c)))
        self.db.commit()
        self.assertNotEqual(service._certificate_manual_base_docx_storage_key(a), service._certificate_manual_base_docx_storage_key(b))
        b_hash = reuse.file_hash(b.storage_key_docx)
        service._rebuild_manual_certificate_word_with_pdf_attachments(self.db, detail=self.detail(first, certificate_id=a.id), certificate=a, actor=self.actor)
        body = '\n'.join(p.text for p in Document(service._certificate_storage_path(a.storage_key_docx)).paragraphs)
        self.assertIn('MANUAL_CONTENT_A', body)
        self.assertNotIn('MANUAL_CONTENT_B', body)
        self.assertEqual(reuse.file_hash(b.storage_key_docx), b_hash)

    def test_download_with_embedded_attachments_keeps_complete_word_and_actors(self):
        _, source = self.base_word(closed=False)
        path = service._certificate_storage_path(source.storage_key_docx)
        doc = Document(path)
        doc.add_paragraph('COMPLETE EMBEDDED MANUAL PAGE')
        doc.save(path)
        source.pdf_attachments_initialized = True
        self.publish(source, 'additional_pages')
        old_key, old_hash = source.storage_key_docx, reuse.file_hash(source.storage_key_docx)
        actors = (source.certified_by_user_id, source.quality_manager_user_id)
        revision = refs.metadata(source)['revision']
        with patch.object(service, '_docx_contains_pdf_attachment_pages', return_value=True), patch.object(service, '_rebuild_generated_certificate_word', side_effect=AssertionError('download must not regenerate technical body')):
            result = service._sync_word_fields_for_download(self.db, certificate=source, path=path)
        self.assertIn('COMPLETE EMBEDDED MANUAL PAGE', '\n'.join(p.text for p in Document(result).paragraphs))
        self.assertEqual(reuse.file_hash(old_key), old_hash)
        self.assertEqual((source.certified_by_user_id, source.quality_manager_user_id), actors)
        self.assertEqual(refs.metadata(source)['revision'], revision)
        self.assertIsNone(refs.choose(self.db, [source])[1])

    def test_legacy_shared_manual_base_is_not_trusted_for_another_record(self):
        first, source = self.base_word(closed=False)
        self.upload(source, text='OWN COMPLETE WORD')
        base = service._certificate_storage_path(service._certificate_manual_base_docx_storage_key(source))
        base.unlink()  # Test-owned file: simulate a legacy installation with only shared base.
        legacy = service._certificate_storage_path('certificati_finali/manual_bases/7013_00_00_26_manual_base.docx')
        legacy.parent.mkdir(parents=True, exist_ok=True)
        wrong = Document();wrong.add_paragraph('ANOTHER RECORD');wrong.save(legacy)
        legacy_hash = reuse.file_hash('certificati_finali/manual_bases/7013_00_00_26_manual_base.docx')
        service._rebuild_manual_certificate_word_with_pdf_attachments(self.db, detail=self.detail(first, certificate_id=source.id), certificate=source, actor=self.actor)
        body = '\n'.join(p.text for p in Document(service._certificate_storage_path(source.storage_key_docx)).paragraphs)
        self.assertIn('OWN COMPLETE WORD', body)
        self.assertNotIn('ANOTHER RECORD', body)
        self.assertEqual(reuse.file_hash('certificati_finali/manual_bases/7013_00_00_26_manual_base.docx'), legacy_hash)

    def test_recovery_preview_reports_without_adopting_or_reassigning(self):
        from app.modules.quarta_taglio.word_lifecycle_audit import preview
        first, source = self.base_word(closed=False)
        self.publish(source)
        old_metadata = dict(refs.metadata(source))
        path = service._certificate_storage_path(source.storage_key_docx)
        doc = Document(path);doc.add_paragraph('UNPUBLISHED');doc.save(path)
        old_key, old_hash = source.storage_key_docx, reuse.file_hash(source.storage_key_docx)
        self.later()
        one, two = preview(self.db), preview(self.db)
        self.assertEqual(one['baseline_hash'], two['baseline_hash'])
        check = next(r for r in one['reference_checks'] if r['id']==source.id)
        self.assertEqual(check['action'], 'review_required_no_automatic_adoption')
        self.assertEqual(check['reason'], 'reference_file_changed')
        self.assertEqual(refs.metadata(source), old_metadata)
        self.assertEqual(source.storage_key_docx, old_key)
        self.assertEqual(reuse.file_hash(old_key), old_hash)
        self.assertFalse(self.db.new or self.db.dirty or self.db.deleted)
        # A legacy missing file must remain a readable diagnostic, not a 500.
        source.word_reference = None
        service._certificate_storage_path(old_key).unlink()
        with self.assertRaises(HTTPException) as caught:
            refs.before_automatic(source)
        self.assertEqual(caught.exception.status_code,409)

    def test_recovery_preview_distinguishes_equivalent_old_path_from_bad_id(self):
        from app.modules.quarta_taglio.word_lifecycle_audit import preview
        _, source = self.base_word(closed=False)
        self.publish(source)
        old_key = source.storage_key_docx
        key = service._certificate_docx_storage_key(source.cod_odp)
        target = service._certificate_storage_path(key)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(service._certificate_storage_path(old_key).read_bytes())
        source.storage_key_docx = key
        self.db.commit()
        report = preview(self.db)
        check = next(r for r in report['reference_checks'] if r['id']==source.id)
        self.assertEqual(check['action'], 'candidate_metadata_path_only')
        self.assertEqual(check['recorded_key'], old_key)
        self.assertFalse(report['identity_issues'])
        self.assertEqual(refs.metadata(source)['storage_key'], old_key)

    def test_field_comparison_does_not_match_an_unrelated_number_in_body(self):
        _, source = self.base_word(closed=False)
        path = service._certificate_storage_path(source.storage_key_docx)
        doc = Document(path);doc.add_paragraph('UNRELATED ORDER HAS 99999');doc.save(path)
        self.assertTrue(service._docx_needs_content_control_update(path, {'QUANTITY_RAW':'99999'}))
        self.assertTrue(service._docx_needs_content_control_update(path, {'DDT_RAW':''}))

    def test_explicit_attachment_edit_keeps_old_revision_path_and_hash_in_history(self):
        first, source = self.base_word(closed=False)
        self.upload(source, text='OWN MANUAL PAGE')
        old = dict(refs.metadata(source))
        detail = self.detail(first, certificate_id=source.id)
        service._rebuild_manual_certificate_word_with_pdf_attachments(self.db,
            detail=detail, certificate=source, actor=self.actor)
        # Internal builder guards the source, but the explicit action is the
        # publication boundary: it must still see the ORIGINAL old revision.
        self.assertEqual(refs.metadata(source), old)
        refs.publish(self.db, source, actor_id=self.actor.id, origin='pdf_attachment')
        self.db.commit()
        current = refs.metadata(source)
        self.assertNotEqual(current['revision'], old['revision'])
        self.assertEqual(current['history'][-1]['storage_key'], old['storage_key'])
        self.assertEqual(current['history'][-1]['template_hash'], old['template_hash'])
        self.assertIsNone(refs.choose(self.db,[source])[1])


for _name in dir(reference_tests.WordReferenceTest):
    if _name.startswith('test_') and _name not in WordLifecycleTest.__dict__:
        setattr(WordLifecycleTest, _name, None)


@unittest.skipUnless(os.environ.get('DDT_TEST_POSTGRES_URL'), 'isolated PostgreSQL test URL not supplied')
class WordLifecyclePostgresTest(WordLifecycleTest):
    setUp = reference_tests.existing.DdtWordReusePostgresTest.setUp

    def test_real_upload_can_interleave_conversion_but_cannot_publish_old_pdf(self):
        self._assert_concurrent_upload_rejected('during_conversion')

    def test_real_upload_after_sync_cannot_authorize_old_input_with_new_state(self):
        self._assert_concurrent_upload_rejected('after_sync')

    def _assert_concurrent_upload_rejected(self, phase):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Event
        from app.core.users.models import User
        from app.modules.quarta_taglio.models import QuartaTaglioFinalCertificate as C
        _, source = self.base_word(closed=False)
        self.publish(source)
        source.conformity_status = 'conforme'
        source.cert_date = datetime.now(timezone.utc)
        self.db.commit()
        certificate_id, actor_id = source.id, self.actor.id
        converting, release = Event(), Event()
        sync = service._sync_word_fields_for_download
        def sync_then_wait(db, *, certificate, path):
            result = sync(db, certificate=certificate, path=path)
            if phase == 'after_sync':
                converting.set()
                if not release.wait(15):
                    raise AssertionError('test upload did not release synchronization gap')
            return result
        def convert(word, pdf):
            self.assertEqual(phase, 'during_conversion', 'old input must be rejected BEFORE conversion')
            converting.set()
            if not release.wait(15):
                raise AssertionError('test upload did not release conversion')
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b'%PDF-1.4 from stable original word')
        def generate():
            with self.factory() as db:
                return service.generate_quarta_taglio_certificate_pdf(db,
                    certificate_id=certificate_id, actor=db.get(User,actor_id))
        def upload():
            with self.factory() as db:
                current = db.get(C,certificate_id)
                doc = Document(service._certificate_storage_path(current.storage_key_docx))
                doc.add_paragraph('A REAL SECOND SESSION UPLOAD')
                stream = BytesIO();doc.save(stream);stream.seek(0)
                return service.upload_quarta_taglio_word_file(db,cod_odp=current.cod_odp,
                    certificate_id=current.id,actor=db.get(User,actor_id),
                    uploaded_file=UploadFile(filename='changed.docx',file=stream))
        with patch.object(service,'_sync_word_fields_for_download',side_effect=sync_then_wait), patch.object(service,'convert_docx_to_pdf',side_effect=convert), patch.object(service,'_build_standard_conformity',return_value=('conforme',[])), ThreadPoolExecutor(max_workers=2) as pool:
            generating = pool.submit(generate)
            try:
                self.assertTrue(converting.wait(10))
                pool.submit(upload).result(timeout=10)
            finally:
                release.set()
            with self.assertRaises(HTTPException) as caught:
                generating.result(timeout=10)
        self.assertEqual(caught.exception.status_code,409)
        self.db.expire_all()
        self.assertNotEqual(self.db.get(C,certificate_id).status,'pdf_final')
        self.assertEqual(self.db.query(Version).count(),0)

    def test_two_real_pdf_requests_publish_one_version_and_remove_unused_output(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from app.core.users.models import User
        _, source = self.base_word(closed=False)
        self.publish(source)
        source.conformity_status = 'conforme'
        source.cert_date = datetime.now(timezone.utc)
        self.db.commit()
        # Stabilize fixture header fields before testing two conversions of the
        # SAME prepared Word (a concurrent header edit is a different scenario).
        with patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            service._sync_word_fields_for_download(self.db, certificate=source,
                path=service._certificate_storage_path(source.storage_key_docx))
        certificate_id, actor_id = source.id, self.actor.id
        converting = Barrier(2, timeout=15)
        outputs = []
        def convert(word, pdf):
            converting.wait()
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b'%PDF-1.4 isolated concurrent conversion')
            outputs.append(pdf)
        def generate():
            with self.factory() as db:
                return service.generate_quarta_taglio_certificate_pdf(db,
                    certificate_id=certificate_id, actor=db.get(User, actor_id))
        with patch.object(service, 'convert_docx_to_pdf', side_effect=convert), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])), ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: generate(), range(2)))
        self.db.expire_all()
        self.assertTrue(all(result.has_pdf for result in results))
        self.assertEqual(self.db.query(Version).filter(Version.status == 'active').count(), 1)
        self.assertEqual(len(outputs), 2)
        self.assertEqual(sum(p.exists() for p in outputs), 1)
