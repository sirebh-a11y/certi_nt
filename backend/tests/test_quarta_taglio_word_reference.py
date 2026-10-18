"""C/D: real documents and an isolated DB; no company connection is permitted."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import os
import unittest

from docx import Document
from fastapi import HTTPException, UploadFile
from sqlalchemy import select

from app.modules.quarta_taglio import service, word_reference as refs, ddt_word_reuse as reuse
from app.modules.quarta_taglio.word_material_update import update_material
from app.modules.quarta_taglio.models import QuartaTaglioFinalCertificate as C
import test_quarta_taglio_ddt_word_reuse as existing


class WordReferenceTest(existing.DdtWordReuseTest):
    def publish(self, source, origin='user_uploaded'):
        refs.publish(self.db, source, actor_id=self.actor.id, origin=origin)
        self.db.commit()

    def edit(self, source, text='SECOND PAGE COMPLETE'):
        old = service._certificate_storage_path(source.storage_key_docx)
        doc = Document(old)
        doc.add_page_break()
        doc.add_paragraph(text)
        key = service._certificate_docx_storage_key(source.cod_odp)
        path = service._certificate_storage_path(key)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        source.storage_key_docx = key
        source.word_source = 'user_uploaded'
        self.publish(source)

    def target(self, item):
        self.db.expire_all()
        return self.db.scalar(select(C).where(C.unit_key == item.certification_unit_key))

    def test_explicit_completion_is_reference_not_latest_created_copy(self):
        _, source = self.base_word(closed=False)
        self.publish(source, 'generated')
        later = self.later()
        self.assertEqual(self.run_worker()['prepared'], 1)
        target = self.target(later)
        self.assertGreater(target.id, source.id)
        self.edit(source)
        self.assertEqual(refs.declared(refs.sources_for(self.db, source.cod_odp, source.cod_f3)).id, source.id)
        self.assertEqual(self.run_worker()['prepared'], 1)
        target = self.target(later)
        self.assertIn('SECOND PAGE COMPLETE', '\n'.join(p.text for p in Document(service._certificate_storage_path(target.storage_key_docx)).paragraphs))
        self.assertEqual(refs.metadata(target)['kind'], 'copy')
        self.assertEqual(refs.metadata(target)['revision'], refs.metadata(source)['revision'])
        self.assertEqual(self.run_worker()['prepared'], 0)

    def test_next_ddt_and_new_workmanship_choose_same_completed_raw(self):
        _, source = self.base_word(closed=False)
        self.publish(source, 'generated')
        self.edit(source)
        later = self.later()
        self.assertEqual(reuse.plan_item(self.db, later)['source_id'], source.id)
        child = C(id=-1, cod_odp=source.cod_odp, cod_f3='001260', certificate_number='NEW')
        chosen = service._previous_word_certificate_for_inheritance(self.db, certificate=child, raw_cod_f3=source.cod_f3)
        self.assertEqual(chosen.id, source.id)

    def test_first_finished_ddt_generation_inherits_embedded_pages_too(self):
        _, source = self.base_word(closed=False)
        self.edit(source, 'RAW LAB PAGE TO INHERIT')
        source.certified_by_user_id = self.actor.id
        self.db.commit()
        finished = self.later(CodF3='001260')
        with patch.object(service, '_word_creation_blockers', return_value=[]), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            response = self.create(finished)
        target = self.db.get(C, response.id)
        self.assertEqual(target.word_source, 'inherited')
        self.assertIn('RAW LAB PAGE TO INHERIT', '\n'.join(p.text for p in Document(service._certificate_storage_path(target.storage_key_docx)).paragraphs))
        self.assertEqual(target.certified_by_user_id, source.certified_by_user_id)
        controls, _ = reuse.word_facts(target.storage_key_docx)
        self.assertEqual(controls['COD_F3_FINISHED'], '001260')
        self.assertEqual(controls['DDT_FINISHED'], finished.ddt_raw)

    def test_legacy_conflict_is_not_silently_selected_for_inheritance(self):
        _, old = self.base_word(closed=False)
        other_item = self.later()
        other = self.cert(other_item, certificate_number=old.certificate_number, cdq_values=deepcopy(old.cdq_values))
        doc = Document(service._certificate_storage_path(old.storage_key_docx))
        doc.add_paragraph('ACTUAL DIFFERENCE')
        doc.save(service._certificate_storage_path(other.storage_key_docx))
        child = C(id=-1, cod_odp=old.cod_odp, cod_f3='001260', certificate_number='NEW')
        self.assertIsNone(service._previous_word_certificate_for_inheritance(self.db, certificate=child, raw_cod_f3=old.cod_f3))
        refs.adopt(self.db, other, self.actor, reuse.file_hash(other.storage_key_docx))
        self.assertEqual(service._previous_word_certificate_for_inheritance(self.db, certificate=child, raw_cod_f3=old.cod_f3).id, other.id)

    def test_closed_pdf_and_word_unchanged_after_reference_edit(self):
        _, closed = self.base_word()
        word_hash, pdf_hash = reuse.file_hash(closed.storage_key_docx), reuse.file_hash(closed.storage_key_pdf)
        later = self.later()
        self.run_worker()
        updated = self.target(later)
        self.edit(updated)
        third = self.later(IdDocumento='300', DDT='2500-20/10/2026')
        self.assertEqual(reuse.plan_item(self.db, third)['source_id'], updated.id)
        self.assertEqual(reuse.file_hash(closed.storage_key_docx), word_hash)
        self.assertEqual(reuse.file_hash(closed.storage_key_pdf), pdf_hash)
        self.assertEqual(closed.status, 'pdf_final')

    def test_manually_changed_open_copy_is_signalled_and_not_overwritten(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        later = self.later()
        self.run_worker()
        target = self.target(later)
        self.edit(target, 'MY MANUAL WORK')
        target_hash = reuse.file_hash(target.storage_key_docx)
        self.edit(source, 'NEW REFERENCE')
        self.assertEqual(reuse.plan_item(self.db, later)['reason'], 'manual_copy_requires_review')
        self.run_worker()
        self.assertEqual(reuse.file_hash(target.storage_key_docx), target_hash)
        self.assertIsNotNone(refs.pending_message(self.db, target))
        row = self.read().items[0]
        row = next(r for r in self.read().items if r.id == later.id)
        self.assertIsNone(row.pdf_action)

    def test_outside_file_change_protects_automatic_copy(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        later = self.later()
        self.run_worker()
        target = self.target(later)
        doc = Document(service._certificate_storage_path(target.storage_key_docx))
        doc.add_paragraph('EDIT OUTSIDE APP')
        doc.save(service._certificate_storage_path(target.storage_key_docx))
        self.edit(source)
        self.assertEqual(reuse.plan_item(self.db, later)['reason'], 'manual_copy_requires_review')

    def test_reference_change_cannot_be_bypassed_by_older_compatible_source(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        later = self.later()
        self.run_worker()
        self.material(cdq='CDQ2', colata='HEAT2')
        source.cdq_values = [dict(v, cdq='CDQ2', colata='HEAT2') for v in source.cdq_values]
        self.publish(source)
        third = self.later(IdDocumento='300')
        self.assertEqual(reuse.plan_item(self.db, third)['reason'], 'no_compatible_word')

    def test_stale_adoption_hash_and_unauthorized_user_are_rejected(self):
        _, source = self.base_word()
        with self.assertRaises(HTTPException) as caught:
            refs.adopt(self.db, source, self.actor, '0' * 64)
        self.assertEqual(caught.exception.status_code, 409)
        with patch.object(service, 'is_quality_area_user', return_value=False):
            with self.assertRaises(HTTPException) as caught:
                refs.adopt(self.db, source, self.actor, reuse.file_hash(source.storage_key_docx))
        self.assertEqual(caught.exception.status_code, 403)

    def test_material_change_has_named_message_and_blocks_pdf_before_conversion(self):
        _, source = self.base_word(closed=False)
        self.material(cdq='CDQ2', colata='HEAT2')
        self.assertIn('CDQ2', refs.material_message(self.db, source))
        with patch.object(service, 'convert_docx_to_pdf', side_effect=AssertionError('must not convert')):
            with self.assertRaises(HTTPException) as caught:
                service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=source.id, actor=self.actor)
        self.assertIn('materiale', str(caught.exception.detail))

    def test_material_refresh_preserves_second_pages_and_all_material_metadata(self):
        _, source = self.base_word(closed=False)
        self.edit(source, 'SECOND PAGE WITH LAB DATA')
        old_key = source.storage_key_docx
        old_hash = reuse.file_hash(old_key)
        self.material(cdq='CDQ2', colata='HEAT2')
        with patch.object(service, '_word_creation_blockers', return_value=[]), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            target = update_material(self.db, source, self.actor, old_hash)
        self.assertEqual(target.id, source.id)
        self.assertEqual({v['cdq'] for v in target.cdq_values}, {'CDQ1','CDQ2'})
        self.assertFalse(refs.material_changed(self.db, target))
        self.assertIn('SECOND PAGE WITH LAB DATA', '\n'.join(p.text for p in Document(service._certificate_storage_path(target.storage_key_docx)).paragraphs))
        self.assertEqual(reuse.file_hash(old_key), old_hash)

    def test_material_refresh_from_closed_pdf_creates_open_reference_without_reopening(self):
        _, source = self.base_word()
        old_word, old_pdf = reuse.file_hash(source.storage_key_docx), reuse.file_hash(source.storage_key_pdf)
        self.material(cdq='CDQ2', colata='HEAT2')
        with patch.object(service, '_word_creation_blockers', return_value=[]), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            target = update_material(self.db, source, self.actor, old_word)
        self.assertNotEqual(target.id, source.id)
        self.assertIsNone(target.ddt)
        self.assertIsNone(target.storage_key_pdf)
        self.assertEqual(source.status, 'pdf_final')
        self.assertEqual(reuse.file_hash(source.storage_key_docx), old_word)
        self.assertEqual(reuse.file_hash(source.storage_key_pdf), old_pdf)

    def test_material_refresh_retains_embedded_images_and_can_render_qa(self):
        import hashlib
        import zipfile
        import shutil
        from PIL import Image, ImageDraw
        first, source = self.base_word(closed=False)
        self.edit(source, 'LABORATORY RESULTS - EXTRA PAGE')
        image = Image.new('RGB', (600, 160), 'white')
        ImageDraw.Draw(image).text((20, 50), 'LAB TEST IMAGE  -  CERTI C/D', fill='black')
        picture = Path(self.files.name, 'lab.png')
        image.save(picture)
        doc = Document(service._certificate_storage_path(source.storage_key_docx))
        from docx.shared import Inches
        doc.add_picture(str(picture), width=Inches(4))
        doc.save(service._certificate_storage_path(source.storage_key_docx))
        self.publish(source)
        old_path = service._certificate_storage_path(source.storage_key_docx)
        def images(path):
            with zipfile.ZipFile(path) as archive:
                return {hashlib.sha256(archive.read(k)).hexdigest() for k in archive.namelist() if k.startswith('word/media/')}
        old_images = images(old_path)
        self.material(cdq='CDQ2', colata='HEAT2')
        from app.modules.quarta_taglio.schemas import QuartaTaglioAggregateValueResponse as Value
        preview = self.detail(first)
        preview.can_create_word = True
        preview.word_creation_blockers = []
        preview.chemistry = [Value(field=field, value=.22, method='weighted', standard_min=0,
                                  standard_max=.4, status='ok') for field in ['Si','Fe','Cu','Mn','Mg','Cr','Zn']]
        preview.properties = [Value(field=field, value=value, method='minimum', standard_min=minimum,
                                    status='ok') for field,value,minimum in [('diametro',50,None),('Rp0.2',300,290),('Rm',350,330),('A%',12,10)]]
        with patch.object(service, '_word_creation_blockers', return_value=[]), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])), patch.object(service, 'get_quarta_taglio_detail', return_value=preview):
            target = update_material(self.db, source, self.actor, reuse.file_hash(source.storage_key_docx))
        updated_path = service._certificate_storage_path(target.storage_key_docx)
        self.assertTrue(old_images.issubset(images(updated_path)))
        if os.environ.get('CERTI_WORD_REFERENCE_QA_DIR'):
            qa = Path(os.environ['CERTI_WORD_REFERENCE_QA_DIR'])
            qa.mkdir(parents=True, exist_ok=True)
            for name, path in [('before', old_path), ('after', updated_path)]:
                shutil.copy2(path, qa / (name + '.docx'))
                service.convert_docx_to_pdf(qa / (name + '.docx'), qa / (name + '.pdf'))
                import fitz
                with fitz.open(qa / (name + '.pdf')) as pdf:
                    self.assertEqual(len(pdf), 2)
                    for index, page in enumerate(pdf):
                        page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2)).save(qa / f'{name}-{index+1}.png')

    def test_reference_routes_reject_bad_hash_and_quality_readonly_user(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from app.core.deps import get_db, get_current_user
        from app.modules.quarta_taglio.router import router
        _, source = self.base_word()
        app = FastAPI()
        app.include_router(router, prefix='/api/quarta-taglio')
        app.dependency_overrides[get_db] = lambda: self.db
        app.dependency_overrides[get_current_user] = lambda: self.actor
        with TestClient(app) as client:
            self.assertEqual(client.post(f'/api/quarta-taglio/certificates/{source.id}/word-reference', json={'source_file_hash':'bad'}).status_code, 422)
            self.assertEqual(client.post(f'/api/quarta-taglio/certificates/{source.id}/word-reference', json={'source_file_hash':'0'*64}).status_code, 409)
            with patch.object(service, 'is_quality_area_user', return_value=False):
                for suffix in ['word-reference','update-material-word','refresh-reference-copy']:
                    self.assertEqual(client.post(f'/api/quarta-taglio/certificates/{source.id}/{suffix}', json={'source_file_hash':reuse.file_hash(source.storage_key_docx)}).status_code, 403)

    def test_explicit_upload_and_regeneration_never_target_another_open_workmanship(self):
        first, source = self.base_word(closed=False)
        preview = self.detail(first)
        preview.can_create_word, preview.word_creation_blockers = True, []
        preview.header['unit_key'] = None
        source.unit_key = None
        other_item = self.later(CodF3='001260')
        other = self.cert(other_item, unit_key=None, certificate_number='9999_00_30/26')
        other_key, other_hash = other.storage_key_docx, reuse.file_hash(other.storage_key_docx)
        self.db.commit()
        content = service._certificate_storage_path(source.storage_key_docx).read_bytes()
        with patch.object(service, 'get_quarta_taglio_detail', return_value=preview):
            uploaded = service.upload_quarta_taglio_word_file(self.db, cod_odp=source.cod_odp,
                certificate_id=source.id, actor=self.actor,
                uploaded_file=UploadFile(filename='completed.docx', file=BytesIO(content)))
        self.assertEqual(uploaded.id, source.id)
        with patch.object(service, 'get_quarta_taglio_detail', return_value=preview), patch.object(service, '_sync_certifiable_unit_register', return_value=[]):
            regenerated = service.create_quarta_taglio_word_draft(self.db, cod_odp=source.cod_odp,
                certificate_id=source.id, actor=self.actor, force_regenerate=True)
        self.assertEqual(regenerated.id, source.id)
        self.assertEqual(other.storage_key_docx, other_key)
        self.assertEqual(reuse.file_hash(other_key), other_hash)

    def test_pdf_closure_rechecks_reference_after_conversion(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        later = self.later()
        self.run_worker()
        target = self.target(later)
        target.conformity_status = 'conforme'
        self.db.commit()
        def convert_and_edit(word, pdf):
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b'%PDF-1.4 fake test output')
            self.edit(source, 'REFERENCE CHANGED DURING CONVERSION')
        with patch.object(service, 'convert_docx_to_pdf', side_effect=convert_and_edit), patch.object(service, '_build_standard_conformity', return_value=('conforme', [])):
            with self.assertRaises(HTTPException) as caught:
                service.generate_quarta_taglio_certificate_pdf(self.db, certificate_id=target.id, actor=self.actor)
        self.assertEqual(caught.exception.status_code, 409)
        self.assertIsNone(target.storage_key_pdf)
        self.assertNotEqual(target.status, 'pdf_final')

    def test_manual_copy_can_be_replaced_after_explicit_confirmation(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        later = self.later()
        self.run_worker()
        target = self.target(later)
        self.edit(target, 'ONLY THIS OLD COPY')
        previous_key = target.storage_key_docx
        previous_hash = reuse.file_hash(previous_key)
        self.edit(source, 'LATEST COMPLETE WORD')
        updated = refs.replace_open_copy(self.db, target, self.actor, previous_hash)
        text = '\n'.join(p.text for p in Document(service._certificate_storage_path(updated.storage_key_docx)).paragraphs)
        self.assertIn('LATEST COMPLETE WORD', text)
        self.assertNotIn('ONLY THIS OLD COPY', text)
        self.assertEqual(reuse.file_hash(previous_key), previous_hash)
        self.assertIsNone(refs.pending_message(self.db, updated))
        self.assertEqual(refs.metadata(updated)['history'][-1]['previous_key'], previous_key)

    def test_copy_of_closed_pdf_cannot_be_replaced(self):
        _, closed = self.base_word()
        with self.assertRaises(HTTPException) as caught:
            refs.replace_open_copy(self.db, closed, self.actor, reuse.file_hash(closed.storage_key_docx))
        self.assertEqual(caught.exception.status_code, 409)

    def test_reference_tamper_does_not_fall_back_to_a_copy(self):
        _, source = self.base_word(closed=False)
        self.publish(source)
        self.later()
        self.run_worker()
        doc = Document(service._certificate_storage_path(source.storage_key_docx))
        doc.add_paragraph('NOT PUBLISHED')
        doc.save(service._certificate_storage_path(source.storage_key_docx))
        third = self.later(IdDocumento='300')
        self.assertEqual(reuse.plan_item(self.db, third)['reason'], 'reference_file_changed')

    def test_unrecognised_material_refresh_retains_original_files_and_metadata(self):
        _, source = self.base_word(closed=False)
        Document().save(service._certificate_storage_path(source.storage_key_docx))
        old_key, old_values = source.storage_key_docx, deepcopy(source.cdq_values)
        old_hash = reuse.file_hash(old_key)
        self.material(cdq='CDQ2', colata='HEAT2')
        with patch.object(service, '_word_creation_blockers', return_value=[]):
            with self.assertRaises(HTTPException):
                update_material(self.db, source, self.actor, old_hash)
        self.db.refresh(source)
        self.assertEqual(source.storage_key_docx, old_key)
        self.assertEqual(source.cdq_values, old_values)
        self.assertEqual(reuse.file_hash(old_key), old_hash)

    def test_multi_material_weight_change_is_detected_even_when_lots_do_not_change(self):
        _, source = self.base_word(closed=False)
        self.material(cdq='CDQ2', colata='HEAT2')
        source.cdq_values = refs.current_materials(self.db, source.cod_odp)
        self.db.commit()
        self.assertFalse(refs.material_changed(self.db, source))
        rows = service._load_quarta_rows_for_detail(self.db, cod_odp=source.cod_odp)
        rows[0].qta_totale *= 2
        self.db.commit()
        self.assertTrue(refs.material_changed(self.db, source))

    def test_new_column_migration_is_idempotent_and_does_not_choose_legacy_reference(self):
        from sqlalchemy import create_engine, text, inspect
        engine = create_engine('sqlite://')
        try:
            with engine.begin() as connection:
                connection.execute(text('CREATE TABLE quarta_taglio_final_certificates (id INTEGER PRIMARY KEY)'))
                self.assertTrue(refs.ensure_schema(connection))
                self.assertFalse(refs.ensure_schema(connection))
                self.assertIn('word_reference', {c['name'] for c in inspect(connection).get_columns('quarta_taglio_final_certificates')})
        finally:
            engine.dispose()


for _name in dir(existing.DdtWordReuseTest):
    if _name.startswith('test_') and _name not in WordReferenceTest.__dict__:
        setattr(WordReferenceTest, _name, None)


@unittest.skipUnless(os.environ.get('DDT_TEST_POSTGRES_URL'), 'isolated PostgreSQL test URL not supplied')
class WordReferencePostgresTest(WordReferenceTest):
    setUp = existing.DdtWordReusePostgresTest.setUp

    def test_concurrent_workers_refresh_same_copy_once(self):
        from concurrent.futures import ThreadPoolExecutor
        _, source = self.base_word(closed=False)
        self.publish(source)
        self.later()
        self.run_worker()
        self.edit(source)
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: reuse.sync_ddt_words(session_factory=self.factory), range(2)))
        self.assertEqual(sum(r['prepared'] for r in results), 1)
        self.assertEqual(sum(r['errors'] for r in results), 0)
