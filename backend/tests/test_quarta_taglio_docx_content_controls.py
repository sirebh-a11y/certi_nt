import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches
from PIL import Image

from app.core.departments.models import Department  # noqa: F401 - ensures SQLAlchemy relationship resolution
from app.core.users.models import User
from app.modules.quarta_taglio.certificate_docx import (
    build_forgialluminio_draft_docx, update_docx_content_controls,
    _append_pdf_attachments, _strip_app_pdf_attachment_pages, PDF_ATTACHMENT_PAGE_MARKER,
)
from app.modules.quarta_taglio.schemas import QuartaTaglioDetailResponse, QuartaTaglioStandardCandidateResponse
from app.modules.quarta_taglio.service import _standard_certificate_material_label
from app.modules.standards.models import NormativeStandard  # noqa: F401 - ensures SQLAlchemy relationship resolution


class QuartaTaglioDocxContentControlTests(unittest.TestCase):
    def test_attachment_removal_restores_original_sections_and_is_idempotent(self):
        def settings(section):
            root = ET.fromstring(section._sectPr.xml)
            # python-docx's section clone intentionally discards Word editing
            # session IDs. Geometry, header/footer relationships and page setup
            # must be identical; rsid attributes are not layout/content.
            for key in list(root.attrib):
                if key.rsplit('}', 1)[-1].startswith('rsid'):
                    del root.attrib[key]
            return ET.tostring(root)
        for legacy in (False, True):
            with self.subTest(legacy=legacy), tempfile.TemporaryDirectory() as folder:
                path, pdf = Path(folder) / 'word.docx', Path(folder) / 'annex.pdf'
                document = Document()
                document.add_paragraph('Original laboratory page')
                section = document.add_section(WD_SECTION.NEW_PAGE)
                section.left_margin = Inches(1.2)
                section.footer.paragraphs[0].text = 'Original signature'
                document.add_paragraph('Second original section')
                document.save(path)
                original_sections = [settings(s) for s in document.sections]
                with fitz.open() as attachment:
                    attachment.new_page().insert_text((70, 100), 'Attachment one')
                    attachment.new_page().insert_text((70, 100), 'Attachment two')
                    attachment.save(pdf)
                user = User(name='Prepared by', email='qa@example.invalid', role='admin')
                _append_pdf_attachments(path, [(pdf, 'annex.pdf')], certified_by=user, quality_manager=user)
                if legacy:
                    document = Document(path)
                    for paragraph in document.paragraphs:
                        if paragraph.text == PDF_ATTACHMENT_PAGE_MARKER:
                            paragraph.text = 'Attachment: legacy.pdf'
                    document.save(path)
                self.assertEqual(len(Document(path).sections), 3)
                _strip_app_pdf_attachment_pages(path)
                clean = Document(path)
                self.assertEqual([settings(s) for s in clean.sections], original_sections)
                self.assertIn('Second original section', [p.text for p in clean.paragraphs])
                self.assertEqual(len(clean.inline_shapes), 0)
                unchanged = path.read_bytes()
                _strip_app_pdf_attachment_pages(path)
                self.assertEqual(path.read_bytes(), unchanged)

    def test_attachment_removal_does_not_cut_later_manual_content(self):
        with tempfile.TemporaryDirectory() as folder:
            path, pdf, image = (Path(folder) / name for name in ('word.docx', 'annex.pdf', 'lab.png'))
            document = Document()
            document.add_paragraph('Original page')
            document.save(path)
            with fitz.open() as attachment:
                attachment.new_page()
                attachment.save(pdf)
            user = User(name='Prepared by', email='qa@example.invalid', role='admin')
            _append_pdf_attachments(path, [(pdf, 'annex.pdf')], certified_by=user, quality_manager=user)
            document = Document(path)
            document.add_page_break()
            document.add_paragraph('Manual page AFTER the app attachment')
            document.add_paragraph('Attachment: this is an ordinary user note')
            document.add_table(rows=1, cols=1).cell(0, 0).text = 'Manual lab results'
            Image.new('RGB', (20, 20), 'green').save(image)
            document.add_picture(str(image), width=Inches(1))
            document.save(path)
            _strip_app_pdf_attachment_pages(path)
            clean = Document(path)
            self.assertIn('Manual page AFTER the app attachment', [p.text for p in clean.paragraphs])
            self.assertIn('Attachment: this is an ordinary user note', [p.text for p in clean.paragraphs])
            self.assertEqual(clean.tables[0].cell(0, 0).text, 'Manual lab results')
            self.assertEqual(len(clean.inline_shapes), 1)
            self.assertNotIn(PDF_ATTACHMENT_PAGE_MARKER, [p.text for p in clean.paragraphs])
            self.assertFalse(clean._body._element.xpath('.//w:br[@w:type="page"]'))

    @staticmethod
    def _minimal_detail() -> QuartaTaglioDetailResponse:
        return QuartaTaglioDetailResponse(
            cod_odp="OLTEST",
            ready=True,
            status_color="green",
            status_message="ok",
            can_create_word=True,
            header={"data_certificato": ""},
            materials=[],
            missing_items=[],
            standard_candidates=[],
            selected_standard=None,
            selected_standard_confirmed=True,
            chemistry=[],
            properties=[],
            notes=[],
            conformity_status="conforme",
            conformity_issues=[],
            esolver_rows=[],
            certifiable_units=[],
        )

    def test_footer_centers_operator_manager_and_signature_in_separate_cells(self) -> None:
        output_path = Path(tempfile.gettempdir()) / "certi_nt_footer_alignment_test.docx"
        user = User(name="Marco Gorza", email="marco@example.test", role="admin")

        build_forgialluminio_draft_docx(
            detail=self._minimal_detail(),
            output_path=output_path,
            draft_number="7000_00_00/26",
            certified_by=user,
            quality_manager=user,
        )

        with zipfile.ZipFile(output_path) as archive:
            footer_xml = next(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
                if name.startswith("word/footer") and name.endswith(".xml")
            )
            document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")

        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        footer_root = ET.fromstring(footer_xml)
        row = footer_root.find(".//w:tr", namespace)
        self.assertIsNotNone(row)
        cells = row.findall("w:tc", namespace)
        self.assertEqual(len(cells), 3)
        grid_widths = [
            int(column.get(f"{{{namespace['w']}}}w"))
            for column in footer_root.findall(".//w:tblGrid/w:gridCol", namespace)
        ]
        self.assertEqual(grid_widths, [5112, 2808, 2304])
        self.assertIn("Operator:", "".join(cells[0].itertext()))
        self.assertIn("Marco Gorza", "".join(cells[0].itertext()))
        self.assertIn("Quality Manager:", "".join(cells[1].itertext()))
        self.assertNotIn("Quality Manager:", "".join(cells[2].itertext()))
        self.assertNotIn("drawing", ET.tostring(cells[1], encoding="unicode"))
        self.assertIn("drawing", ET.tostring(cells[2], encoding="unicode"))
        self.assertEqual(
            [cell.find("w:tcPr/w:vAlign", namespace).get(f"{{{namespace['w']}}}val") for cell in cells],
            ["center", "center", "center"],
        )
        self.assertIn('w:footer="173"', document_xml)
        self.assertIn('w:bottom="792"', document_xml)
        self.assertIn('cx="512064"', footer_xml)

    def test_generated_certificate_alloy_excludes_standard_variant(self) -> None:
        standard = NormativeStandard(
            code="2024-sigma",
            lega_base="2024",
            lega_designazione="2024 Sigma",
            variante_lega="Sigma",
        )
        detail = QuartaTaglioDetailResponse(
            cod_odp="OLTEST",
            ready=True,
            status_color="green",
            status_message="ok",
            can_create_word=True,
            header={"data_certificato": ""},
            materials=[],
            missing_items=[],
            standard_candidates=[],
            selected_standard=QuartaTaglioStandardCandidateResponse(
                id=1,
                code=standard.code,
                label="2024 Sigma · EN 755-2",
                lega_base=standard.lega_base,
                lega_designazione=standard.lega_designazione,
                variante_lega=standard.variante_lega,
                norma="EN 755-2",
                certificate_material_label=_standard_certificate_material_label(standard),
                confidence="confermata",
                score=999,
            ),
            selected_standard_confirmed=True,
            chemistry=[],
            properties=[],
            notes=[],
            conformity_status="conforme",
            conformity_issues=[],
            esolver_rows=[],
            certifiable_units=[],
        )
        user = User(name="System Admin", email="system@example.test", role="admin")
        output_path = Path(tempfile.gettempdir()) / "certi_nt_standard_variant_certificate_test.docx"

        build_forgialluminio_draft_docx(
            detail=detail,
            output_path=output_path,
            draft_number="7000_00_00/26",
            certified_by=user,
            quality_manager=user,
        )

        with zipfile.ZipFile(output_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")

        self.assertIn("EN AW 2024", document_xml)
        self.assertNotIn("Sigma", document_xml)

    def test_header_dynamic_fields_are_word_content_controls(self) -> None:
        detail = QuartaTaglioDetailResponse(
            cod_odp="OLTEST",
            ready=True,
            status_color="green",
            status_message="ok",
            can_create_word=True,
            header={
                "data_certificato": "",
                "cliente": "CLIENTE TEST",
                "ordine_cliente": "ORD123",
                "conferma_ordine": "CDO456",
                "codice_f3_raw": "100",
                "descrizione_raw": "RAW DESC",
                "ddt_raw": "",
                "quantita_raw": "10",
                "codice_f3_finished": "200",
                "descrizione_finished": "",
                "ddt_finished": "",
                "quantita_finished": "",
            },
            materials=[],
            missing_items=[],
            standard_candidates=[],
            selected_standard=None,
            selected_standard_confirmed=True,
            chemistry=[],
            properties=[],
            notes=[],
            conformity_status="conforme",
            conformity_issues=[],
            esolver_rows=[],
            certifiable_units=[],
        )
        user = User(name="System Admin", email="system@example.test", role="admin")
        output_path = Path(tempfile.gettempdir()) / "certi_nt_content_controls_test.docx"

        build_forgialluminio_draft_docx(
            detail=detail,
            output_path=output_path,
            draft_number="7000_00_00/26",
            certified_by=user,
            quality_manager=user,
        )

        with zipfile.ZipFile(output_path) as archive:
            header_xml = "".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
                if name.startswith("word/header")
            )

        expected_tags = {
            "CERT_NUMBER",
            "CERT_DATE",
            "PURCHASER",
            "ORDER_CLIENT",
            "CONFIRM_ORDER",
            "COD_F3_RAW",
            "RAW_DESCRIPTION",
            "DDT_RAW",
            "QUANTITY_RAW",
            "COD_F3_FINISHED",
            "FINISHED_DESCRIPTION",
            "DDT_FINISHED",
            "QUANTITY_FINISHED",
        }
        for tag in expected_tags:
            self.assertIn(f'w:val="{tag}"', header_xml)
        self.assertEqual(header_xml.count("<w:sdt>"), len(expected_tags))
        self.assertEqual(header_xml.count("Delivery Note:"), 2)
        self.assertEqual(header_xml.count("(D.d.T.)"), 2)
        self.assertNotIn("T.D.:", header_xml)
        self.assertNotIn("(D.d.T.):", header_xml)
        self.assertIn("Quantity:", header_xml)
        self.assertIn("Quantit", header_xml)

    def test_update_content_controls_changes_only_tagged_values(self) -> None:
        detail = QuartaTaglioDetailResponse(
            cod_odp="OLTEST",
            ready=True,
            status_color="green",
            status_message="ok",
            can_create_word=True,
            header={
                "data_certificato": "",
                "cliente": "CLIENTE TEST",
                "ordine_cliente": "ORD123",
                "conferma_ordine": "CDO456",
                "codice_f3_raw": "100",
                "descrizione_raw": "RAW DESC",
                "ddt_raw": "",
                "quantita_raw": "10",
                "codice_f3_finished": "",
                "descrizione_finished": "",
                "ddt_finished": "",
                "quantita_finished": "",
            },
            materials=[],
            missing_items=[],
            standard_candidates=[],
            selected_standard=None,
            selected_standard_confirmed=True,
            chemistry=[],
            properties=[],
            notes=[],
            conformity_status="conforme",
            conformity_issues=[],
            esolver_rows=[],
            certifiable_units=[],
        )
        user = User(name="System Admin", email="system@example.test", role="admin")
        source_path = Path(tempfile.gettempdir()) / "certi_nt_content_controls_source.docx"
        output_path = Path(tempfile.gettempdir()) / "certi_nt_content_controls_updated.docx"
        build_forgialluminio_draft_docx(
            detail=detail,
            output_path=source_path,
            draft_number="7000_00_00/26",
            certified_by=user,
            quality_manager=user,
        )

        updated, missing = update_docx_content_controls(
            source_path,
            output_path,
            {
                "CERT_DATE": "19/05/2026",
                "DDT_RAW": "1133-19/05/2026",
                "DDT_FINISHED": "2026/DDT-000123456789-19/05/2026",
                "COD_F3_FINISHED": "605001860",
            },
        )

        self.assertIn("CERT_DATE", updated)
        self.assertIn("DDT_RAW", updated)
        self.assertIn("DDT_FINISHED", updated)
        self.assertIn("COD_F3_FINISHED", updated)
        self.assertEqual(missing, [])
        with zipfile.ZipFile(output_path) as archive:
            header_xml = "".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in archive.namelist()
                if name.startswith("word/header")
            )
            for name in archive.namelist():
                if name.startswith("word/") and name.endswith(".xml"):
                    ET.fromstring(archive.read(name))
        self.assertIn("19/05/2026", header_xml)
        self.assertIn("1133-19/05/2026", header_xml)
        self.assertIn("2026/DDT-000123456789-19/05/2026", header_xml)
        self.assertEqual(header_xml.count("Delivery Note:"), 2)
        self.assertNotIn("T.D.:", header_xml)
        self.assertIn("605001860", header_xml)
        self.assertIn("<w:text/>", header_xml)
        self.assertIn("<w:sdtContent>", header_xml)


if __name__ == "__main__":
    unittest.main()
